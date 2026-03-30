from __future__ import annotations

import base64
import csv
import io
import json
import logging
import math
import os
import re
import uuid
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Generator, List, Optional, Union
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlparse
from urllib.request import Request, urlopen

from fastapi import Depends, FastAPI, HTTPException
from minio import Minio
from pydantic import BaseModel, Field
from sqlalchemy import DateTime, ForeignKey, Integer, String, Text, create_engine, select
from sqlalchemy.engine import make_url
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column, relationship, sessionmaker

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover
    DocxDocument = None

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover
    PdfReader = None

def now_utc() -> datetime:
    return datetime.now(timezone.utc)


SERVICE_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DATA_ROOT = SERVICE_ROOT / "data"
DEFAULT_DATABASE_PATH = DEFAULT_DATA_ROOT / "knowledge-service.db"
DEFAULT_STORAGE_ROOT = DEFAULT_DATA_ROOT / "storage"


def ensure_sqlite_parent(database_url: str) -> None:
    try:
        url = make_url(database_url)
    except Exception:
        return
    if not url.drivername.startswith("sqlite") or url.database in (None, "", ":memory:"):
        return
    database_path = Path(url.database)
    if not database_path.is_absolute():
        database_path = (Path.cwd() / database_path).resolve()
    database_path.parent.mkdir(parents=True, exist_ok=True)


DATABASE_URL = os.getenv("LYNXUS_KNOWLEDGE_DATABASE_URL", f"sqlite+pysqlite:///{DEFAULT_DATABASE_PATH}")
STORAGE_MODE = os.getenv("LYNXUS_KNOWLEDGE_STORAGE_MODE", "filesystem").lower()
STORAGE_ROOT = Path(os.getenv("LYNXUS_KNOWLEDGE_STORAGE_ROOT", str(DEFAULT_STORAGE_ROOT)))
MINIO_ENDPOINT = os.getenv("LYNXUS_MINIO_ENDPOINT", "localhost:9000").replace("http://", "").replace("https://", "")
MINIO_ACCESS_KEY = os.getenv("LYNXUS_MINIO_ACCESS_KEY", "minioadmin")
MINIO_SECRET_KEY = os.getenv("LYNXUS_MINIO_SECRET_KEY", "minioadmin")
MINIO_SECURE = os.getenv("LYNXUS_MINIO_ENDPOINT", "http://localhost:9000").startswith("https://")
MINIO_BUCKET = os.getenv("LYNXUS_KNOWLEDGE_MINIO_BUCKET", "lynxus-knowledge")
URL_IMPORT_TIMEOUT_SECONDS = float(os.getenv("LYNXUS_KNOWLEDGE_URL_IMPORT_TIMEOUT_SECONDS", "15"))
URL_IMPORT_USER_AGENT = os.getenv(
    "LYNXUS_KNOWLEDGE_URL_IMPORT_USER_AGENT",
    "LynxusKnowledgeService/2.0 (+https://lynxus.local)",
)
OPENSEARCH_URL = os.getenv("LYNXUS_OPENSEARCH_URL", "").rstrip("/")
OPENSEARCH_USERNAME = os.getenv("LYNXUS_OPENSEARCH_USERNAME", "")
OPENSEARCH_PASSWORD = os.getenv("LYNXUS_OPENSEARCH_PASSWORD", "")
OPENSEARCH_INDEX_PREFIX = os.getenv("LYNXUS_OPENSEARCH_INDEX_PREFIX", "lynxus-knowledge")
OPENSEARCH_TIMEOUT_SECONDS = float(os.getenv("LYNXUS_OPENSEARCH_TIMEOUT_SECONDS", "10"))
DEFAULT_SNAPSHOT_RETRIEVAL_MODE = os.getenv("LYNXUS_KNOWLEDGE_DEFAULT_RETRIEVAL_MODE", "HYBRID").upper()
DEFAULT_SNAPSHOT_RETRIEVAL_BACKEND = os.getenv("LYNXUS_KNOWLEDGE_DEFAULT_RETRIEVAL_BACKEND", "OPENSEARCH").upper()

ensure_sqlite_parent(DATABASE_URL)
if STORAGE_MODE == "filesystem":
    STORAGE_ROOT.mkdir(parents=True, exist_ok=True)

engine = create_engine(DATABASE_URL, future=True)
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, expire_on_commit=False)
logger = logging.getLogger(__name__)

SUPPORTED_FILE_TYPES = ["pdf", "docx", "md", "txt", "html", "csv"]
SNAPSHOT_RETRIEVAL_MODES = {"LEXICAL", "VECTOR", "HYBRID"}
MAX_RERANK_CANDIDATES = 20


class Base(DeclarativeBase):
    pass


class UploadSessionRecord(Base):
    __tablename__ = "knowledge_upload_session"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    knowledge_base_id: Mapped[str] = mapped_column(String(64), index=True)
    status: Mapped[str] = mapped_column(String(32), default="OPEN")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)


class KnowledgeFileRecord(Base):
    __tablename__ = "knowledge_file"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    knowledge_base_id: Mapped[str] = mapped_column(String(64), index=True)
    upload_session_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_upload_session.id"))
    file_name: Mapped[str] = mapped_column(String(255))
    content_type: Mapped[str] = mapped_column(String(128))
    object_key: Mapped[str] = mapped_column(String(255))
    size_bytes: Mapped[int] = mapped_column(Integer)
    status: Mapped[str] = mapped_column(String(32), default="UPLOADED")
    error_message: Mapped[Optional[str]] = mapped_column(Text, nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)

    import_jobs: Mapped[List["KnowledgeImportJobRecord"]] = relationship(back_populates="knowledge_file")


class KnowledgeImportJobRecord(Base):
    __tablename__ = "knowledge_import_job"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    knowledge_base_id: Mapped[str] = mapped_column(String(64), index=True)
    file_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_file.id"))
    status: Mapped[str] = mapped_column(String(32), default="PENDING")
    failure_reason: Mapped[Optional[str]] = mapped_column(Text, nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    completed_at: Mapped[Optional[datetime]] = mapped_column(DateTime(timezone=True), nullable=True)

    knowledge_file: Mapped[KnowledgeFileRecord] = relationship(back_populates="import_jobs")


class KnowledgeDocumentRecord(Base):
    __tablename__ = "knowledge_document"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    knowledge_base_id: Mapped[str] = mapped_column(String(64), index=True)
    file_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_file.id"))
    title: Mapped[str] = mapped_column(String(255))
    source_uri: Mapped[str] = mapped_column(String(255))
    document_type: Mapped[str] = mapped_column(String(32))
    status: Mapped[str] = mapped_column(String(32), default="READY")
    body_text: Mapped[str] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)

    chunks: Mapped[List["KnowledgeChunkRecord"]] = relationship(back_populates="document")


class KnowledgeChunkRecord(Base):
    __tablename__ = "knowledge_chunk"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    knowledge_base_id: Mapped[str] = mapped_column(String(64), index=True)
    document_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_document.id"))
    chunk_index: Mapped[int] = mapped_column(Integer)
    title: Mapped[str] = mapped_column(String(255))
    heading_path: Mapped[str] = mapped_column(String(512), default="")
    source_uri: Mapped[str] = mapped_column(String(255), default="")
    page_number: Mapped[Optional[int]] = mapped_column(Integer, nullable=True)
    content: Mapped[str] = mapped_column(Text)
    token_count: Mapped[int] = mapped_column(Integer)
    normalized_title: Mapped[str] = mapped_column(Text)
    normalized_heading_path: Mapped[str] = mapped_column(Text)
    normalized_content: Mapped[str] = mapped_column(Text)
    normalized_source_uri: Mapped[str] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)

    document: Mapped[KnowledgeDocumentRecord] = relationship(back_populates="chunks")


class IndexSnapshotRecord(Base):
    __tablename__ = "knowledge_index_snapshot"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    knowledge_base_id: Mapped[str] = mapped_column(String(64), index=True)
    retrieval_backend: Mapped[str] = mapped_column(String(64), default=DEFAULT_SNAPSHOT_RETRIEVAL_BACKEND)
    retrieval_mode: Mapped[str] = mapped_column(String(32), default=DEFAULT_SNAPSHOT_RETRIEVAL_MODE)
    status: Mapped[str] = mapped_column(String(32), default="PENDING")
    document_count: Mapped[int] = mapped_column(Integer, default=0)
    chunk_count: Mapped[int] = mapped_column(Integer, default=0)
    failure_reason: Mapped[Optional[str]] = mapped_column(Text, nullable=True)
    built_at: Mapped[Optional[datetime]] = mapped_column(DateTime(timezone=True), nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)

    snapshot_chunks: Mapped[List["IndexSnapshotChunkRecord"]] = relationship(back_populates="snapshot")
    selected_documents: Mapped[List["SnapshotDocumentSelectionRecord"]] = relationship(back_populates="snapshot")


class IndexSnapshotChunkRecord(Base):
    __tablename__ = "knowledge_index_snapshot_chunk"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    snapshot_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_index_snapshot.id"), index=True)
    chunk_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_chunk.id"))
    document_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_document.id"))

    snapshot: Mapped[IndexSnapshotRecord] = relationship(back_populates="snapshot_chunks")


class SnapshotDocumentSelectionRecord(Base):
    __tablename__ = "knowledge_snapshot_document_selection"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    snapshot_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_index_snapshot.id"), index=True)
    document_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_document.id"))

    snapshot: Mapped[IndexSnapshotRecord] = relationship(back_populates="selected_documents")


class CreateUploadSessionRequest(BaseModel):
    knowledgeBaseId: str


class UploadSessionResponse(BaseModel):
    id: str
    knowledgeBaseId: str
    status: str
    acceptedTypes: List[str]


class CompleteUploadRequest(BaseModel):
    knowledgeBaseId: str
    uploadSessionId: str
    fileName: str
    contentType: str
    contentBase64: str


class CreateUrlImportRequest(BaseModel):
    knowledgeBaseId: str
    url: str
    title: Optional[str] = None


class KnowledgeFileResponse(BaseModel):
    id: str
    knowledgeBaseId: str
    uploadSessionId: str
    fileName: str
    contentType: str
    sizeBytes: int
    status: str
    errorMessage: Optional[str] = None
    createdAt: datetime
    updatedAt: datetime


class ImportJobResponse(BaseModel):
    id: str
    knowledgeBaseId: str
    fileId: str
    status: str
    failureReason: Optional[str] = None
    createdAt: datetime
    updatedAt: datetime
    completedAt: Optional[datetime] = None


class DocumentResponse(BaseModel):
    id: str
    knowledgeBaseId: str
    fileId: str
    title: str
    sourceUri: str
    documentType: str
    status: str
    chunkCount: int
    createdAt: datetime
    updatedAt: datetime


class CreateIndexSnapshotRequest(BaseModel):
    knowledgeBaseId: Optional[str] = None
    documentIds: List[str] = Field(default_factory=list)
    retrievalMode: str = Field(default=DEFAULT_SNAPSHOT_RETRIEVAL_MODE)


class IndexSnapshotResponse(BaseModel):
    id: str
    knowledgeBaseId: str
    retrievalBackend: str
    retrievalMode: str
    status: str
    documentCount: int
    chunkCount: int
    failureReason: Optional[str] = None
    builtAt: Optional[datetime] = None
    createdAt: datetime
    updatedAt: datetime


class RetrieveRequest(BaseModel):
    indexSnapshotId: str
    query: str
    topK: int = 5
    minScore: float = 0.1
    retrievalMode: Optional[str] = None


class RetrieveHit(BaseModel):
    chunkId: str
    documentId: str
    documentTitle: str
    sourceUri: str
    snippet: str
    score: float
    pageNumber: Optional[int] = None
    headingPath: str = ""


class RetrieveResponse(BaseModel):
    hits: List[RetrieveHit]
    lowConfidence: bool


@dataclass
class ParsedDocument:
    title: str
    document_type: str
    body_text: str
    segments: List["ParsedSegment"]


@dataclass
class ParsedSegment:
    title: str
    heading_path: str
    content: str
    page_number: Optional[int] = None


class Storage:
    def __init__(self) -> None:
        self.client = None
        if STORAGE_MODE == "minio":
            self.client = Minio(
                MINIO_ENDPOINT,
                access_key=MINIO_ACCESS_KEY,
                secret_key=MINIO_SECRET_KEY,
                secure=MINIO_SECURE,
            )
            self._ensure_bucket()
        else:
            STORAGE_ROOT.mkdir(parents=True, exist_ok=True)

    def _ensure_bucket(self) -> None:
        assert self.client is not None
        if not self.client.bucket_exists(MINIO_BUCKET):
            self.client.make_bucket(MINIO_BUCKET)

    def put_bytes(self, object_key: str, payload: bytes, content_type: str) -> None:
        if self.client is None:
            target = STORAGE_ROOT / object_key
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(payload)
            return
        stream = io.BytesIO(payload)
        self.client.put_object(
            MINIO_BUCKET,
            object_key,
            stream,
            length=len(payload),
            content_type=content_type or "application/octet-stream",
        )

    def get_bytes(self, object_key: str) -> bytes:
        if self.client is None:
            return (STORAGE_ROOT / object_key).read_bytes()
        response = self.client.get_object(MINIO_BUCKET, object_key)
        try:
            return response.read()
        finally:
            response.close()
            response.release_conn()


class OpenSearchClient:
    def __init__(self) -> None:
        self.base_url = OPENSEARCH_URL
        self.username = OPENSEARCH_USERNAME
        self.password = OPENSEARCH_PASSWORD
        self.index_prefix = OPENSEARCH_INDEX_PREFIX.strip() or "lynxus-knowledge"
        self.timeout_seconds = OPENSEARCH_TIMEOUT_SECONDS

    @property
    def enabled(self) -> bool:
        return bool(self.base_url)

    def snapshot_index_name(self, snapshot_id: str) -> str:
        sanitized = re.sub(r"[^a-z0-9-]+", "-", snapshot_id.lower()).strip("-")
        return f"{self.index_prefix}-{sanitized or 'snapshot'}"

    def request(self, method: str, path: str, payload: Optional[Union[dict, list, str]] = None, content_type: str = "application/json") -> dict:
        if not self.enabled:
            raise ValueError("opensearch is not configured")
        headers = {"Accept": "application/json"}
        data: Optional[bytes] = None
        if payload is not None:
            if content_type == "application/x-ndjson" and isinstance(payload, str):
                data = payload.encode("utf-8")
            else:
                data = json.dumps(payload).encode("utf-8")
            headers["Content-Type"] = content_type
        request = Request(f"{self.base_url}{path}", data=data, method=method.upper(), headers=headers)
        if self.username or self.password:
            token = base64.b64encode(f"{self.username}:{self.password}".encode("utf-8")).decode("ascii")
            request.add_header("Authorization", f"Basic {token}")
        try:
            with urlopen(request, timeout=self.timeout_seconds) as response:
                body = response.read().decode("utf-8", errors="ignore").strip()
                return json.loads(body) if body else {}
        except HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="ignore")
            if exc.code == 404:
                raise KeyError(detail or "not found") from exc
            raise ValueError(f"opensearch {method.upper()} {path} failed: {exc.code} {detail}") from exc
        except URLError as exc:
            raise ValueError(f"opensearch request failed: {exc.reason}") from exc

    def exists_index(self, index_name: str) -> bool:
        try:
            self.request("HEAD", f"/{quote(index_name)}")
            return True
        except KeyError:
            return False

    def delete_index(self, index_name: str) -> None:
        if not self.exists_index(index_name):
            return
        self.request("DELETE", f"/{quote(index_name)}")

    def ensure_index(self, index_name: str) -> None:
        if self.exists_index(index_name):
            return
        self.request(
            "PUT",
            f"/{quote(index_name)}",
            {
                "settings": {
                    "index": {"number_of_shards": 1, "number_of_replicas": 0},
                    "analysis": {
                        "analyzer": {
                            "lynxus_text": {
                                "type": "custom",
                                "tokenizer": "standard",
                                "filter": ["lowercase"],
                            }
                        }
                    },
                },
                "mappings": {
                    "properties": {
                        "snapshot_id": {"type": "keyword"},
                        "knowledge_base_id": {"type": "keyword"},
                        "document_id": {"type": "keyword"},
                        "chunk_id": {"type": "keyword"},
                        "chunk_index": {"type": "integer"},
                        "page_number": {"type": "integer"},
                        "title": {"type": "text", "analyzer": "lynxus_text"},
                        "heading_path": {"type": "text", "analyzer": "lynxus_text"},
                        "source_uri": {"type": "text", "analyzer": "lynxus_text"},
                        "content": {"type": "text", "analyzer": "lynxus_text"},
                        "combined_text": {"type": "text", "analyzer": "lynxus_text"},
                    }
                },
            },
        )

    def bulk_index_chunks(self, snapshot_id: str, knowledge_base_id: str, chunks: List[KnowledgeChunkRecord]) -> None:
        index_name = self.snapshot_index_name(snapshot_id)
        self.delete_index(index_name)
        self.ensure_index(index_name)
        lines: List[str] = []
        for chunk in chunks:
            lines.append(json.dumps({"index": {"_index": index_name, "_id": chunk.id}}, ensure_ascii=False))
            lines.append(
                json.dumps(
                    {
                        "snapshot_id": snapshot_id,
                        "knowledge_base_id": knowledge_base_id,
                        "document_id": chunk.document_id,
                        "chunk_id": chunk.id,
                        "chunk_index": chunk.chunk_index,
                        "page_number": chunk.page_number,
                        "title": chunk.title,
                        "heading_path": chunk.heading_path,
                        "source_uri": chunk.source_uri,
                        "content": chunk.content,
                        "combined_text": "\n".join(
                            item for item in [chunk.title, chunk.heading_path, chunk.source_uri, chunk.content] if item
                        ),
                    },
                    ensure_ascii=False,
                )
            )
        payload = "\n".join(lines) + "\n"
        response = self.request("POST", "/_bulk?refresh=true", payload, content_type="application/x-ndjson")
        if response.get("errors"):
            raise ValueError("opensearch bulk indexing reported item errors")

    def search_chunk_ids(self, snapshot_id: str, query: str, retrieval_mode: str, size: int) -> List[str]:
        index_name = self.snapshot_index_name(snapshot_id)
        normalized_mode = normalize_retrieval_mode(retrieval_mode)
        fuzziness = "AUTO" if normalized_mode != "VECTOR" else "0"
        result = self.request(
            "POST",
            f"/{quote(index_name)}/_search",
            {
                "size": size,
                "_source": ["chunk_id"],
                "query": {
                    "bool": {
                        "must": [
                            {
                                "multi_match": {
                                    "query": query,
                                    "fields": [
                                        "title^4",
                                        "heading_path^2",
                                        "content",
                                        "source_uri^0.5",
                                        "combined_text^1.5",
                                    ],
                                    "type": "best_fields",
                                    "operator": "or",
                                    "fuzziness": fuzziness,
                                }
                            }
                        ],
                        "filter": [{"term": {"snapshot_id": snapshot_id}}],
                    }
                },
            },
        )
        hits = result.get("hits", {}).get("hits", [])
        return [hit.get("_source", {}).get("chunk_id") or hit.get("_id") for hit in hits if hit]


storage = Storage()
opensearch = OpenSearchClient()
app = FastAPI(title="Lynxus Knowledge Service", version="2.0.0")


def get_db() -> Generator[Session, None, None]:
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().lower()


def tokenize(value: str) -> List[str]:
    normalized = normalize_text(value)
    if not normalized:
        return []
    tokens = re.findall(r"[a-z0-9_]+|[\u4e00-\u9fff]", normalized)
    if len(tokens) <= 1:
        bigrams = [normalized[index : index + 2] for index in range(max(0, len(normalized) - 1))]
        return bigrams or [normalized]
    return tokens


def estimate_tokens(value: str) -> int:
    return max(1, len(tokenize(value)))


def split_sentences(value: str) -> List[str]:
    parts = re.split(r"(?<=[。！？.!?])\s+|\n+", value)
    return [part.strip() for part in parts if part.strip()]


def detect_document_type(file_name: str, content_type: Optional[str] = None) -> str:
    suffix = Path(file_name).suffix.lower().lstrip(".")
    if suffix in SUPPORTED_FILE_TYPES:
        return suffix
    normalized_content_type = (content_type or "").lower()
    if "markdown" in normalized_content_type:
        return "md"
    if "html" in normalized_content_type:
        return "html"
    if "csv" in normalized_content_type:
        return "csv"
    if "plain" in normalized_content_type or "text/" in normalized_content_type:
        return "txt"
    if "pdf" in normalized_content_type:
        return "pdf"
    if "wordprocessingml" in normalized_content_type or "msword" in normalized_content_type:
        return "docx"
    return suffix or "txt"


def ensure_supported_document_type(file_name: str, content_type: Optional[str]) -> str:
    document_type = detect_document_type(file_name, content_type)
    if document_type not in SUPPORTED_FILE_TYPES:
        raise ValueError(f"unsupported file type: {document_type}")
    return document_type


def html_to_text_payload(raw_html: str) -> tuple[str, str]:
    html = re.sub(r"<script.*?>.*?</script>", " ", raw_html, flags=re.IGNORECASE | re.DOTALL)
    html = re.sub(r"<style.*?>.*?</style>", " ", html, flags=re.IGNORECASE | re.DOTALL)
    html = re.sub(r"(?i)<br\\s*/?>", "\n", html)
    html = re.sub(r"(?i)</p>", "\n\n", html)
    html = re.sub(r"(?i)</div>", "\n", html)
    html = re.sub(r"(?i)</li>", "\n", html)
    html = re.sub(r"(?i)</tr>", "\n", html)
    for level in range(6, 0, -1):
        pattern = rf"(?is)<h{level}[^>]*>(.*?)</h{level}>"
        html = re.sub(pattern, lambda match: f"\n\n{'#' * level} {strip_html_inline(match.group(1))}\n\n", html)
    html = re.sub(r"(?is)<title[^>]*>(.*?)</title>", lambda match: f"\n\n# {strip_html_inline(match.group(1))}\n\n", html)
    text = strip_html_preserving_breaks(html)
    title_match = re.search(r"(?is)<title[^>]*>(.*?)</title>", raw_html)
    if not title_match:
        title_match = re.search(r"(?is)<h1[^>]*>(.*?)</h1>", raw_html)
    title = strip_html_inline(title_match.group(1)) if title_match else ""
    return title, text


def strip_html_inline(value: str) -> str:
    value = re.sub(r"(?is)<[^>]+>", " ", value)
    value = value.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
    return re.sub(r"\s+", " ", value).strip()


def strip_html_preserving_breaks(value: str) -> str:
    value = re.sub(r"(?is)<[^>]+>", " ", value)
    value = value.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
    value = re.sub(r"\r\n?", "\n", value)
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r" *\n *", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def extract_pdf_text(payload: bytes) -> str:
    if PdfReader is None:
        raise ValueError("pdf parsing dependency is unavailable")
    reader = PdfReader(io.BytesIO(payload))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()


def extract_text(file_name: str, payload: bytes, content_type: Optional[str]) -> tuple[str, str]:
    document_type = ensure_supported_document_type(file_name, content_type)
    if document_type in {"txt", "md"}:
        return document_type, payload.decode("utf-8", errors="ignore")
    if document_type == "html":
        _, text = html_to_text_payload(payload.decode("utf-8", errors="ignore"))
        return document_type, text
    if document_type == "csv":
        return document_type, payload.decode("utf-8", errors="ignore")
    if document_type == "pdf":
        return document_type, extract_pdf_text(payload)
    if document_type == "docx":
        if DocxDocument is None:
            raise ValueError("docx parsing dependency is unavailable")
        document = DocxDocument(io.BytesIO(payload))
        return document_type, "\n\n".join(paragraph.text for paragraph in document.paragraphs)
    raise ValueError(f"unsupported file type: {document_type}")


def parse_document(file_record: KnowledgeFileRecord, payload: bytes) -> ParsedDocument:
    document_type = ensure_supported_document_type(file_record.file_name, file_record.content_type)
    source_uri = infer_source_uri(file_record)
    if document_type == "html":
        raw_html = payload.decode("utf-8", errors="ignore")
        title, text = html_to_text_payload(raw_html)
        segments = parse_markdown_segments(file_record.file_name, text)
        return ParsedDocument(
            title=title or Path(file_record.file_name).stem,
            document_type=document_type,
            body_text=text,
            segments=segments,
        )
    text_type, text = extract_text(file_record.file_name, payload, file_record.content_type)
    segments = parse_segments(file_record.file_name, text, source_uri)
    return ParsedDocument(
        title=Path(file_record.file_name).stem,
        document_type=text_type,
        body_text=text,
        segments=segments,
    )


def parse_segments(file_name: str, text: str, source_uri: str) -> List[ParsedSegment]:
    document_type = detect_document_type(file_name)
    if document_type == "md":
        return parse_markdown_segments(file_name, text)
    if document_type == "csv":
        return parse_csv_segments(file_name, text)
    if document_type == "html":
        return parse_markdown_segments(file_name, text)
    return parse_plain_segments(file_name, text, source_uri)


def parse_markdown_segments(file_name: str, text: str) -> List[ParsedSegment]:
    heading_stack: List[str] = []
    segments: List[ParsedSegment] = []
    current: List[str] = []

    def flush() -> None:
        content = "\n".join(current).strip()
        if not content:
            current.clear()
            return
        title = heading_stack[-1] if heading_stack else Path(file_name).stem
        segments.extend(faq_or_plain_segments(title, " / ".join(heading_stack), content))
        current.clear()

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if is_markdown_table_separator(line):
            current.append(line)
            continue
        if not line.strip():
            current.append("")
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            flush()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            heading_stack[:] = heading_stack[: level - 1]
            heading_stack.append(title)
            continue
        current.append(line)
    flush()
    return segments or [ParsedSegment(Path(file_name).stem, "", text.strip())]


def is_markdown_table_separator(line: str) -> bool:
    candidate = line.strip().replace("|", "").replace(":", "").replace("-", "")
    return not candidate and "|" in line and "-" in line


def parse_csv_segments(file_name: str, text: str) -> List[ParsedSegment]:
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return []
    header = [cell.strip() or f"column_{index + 1}" for index, cell in enumerate(rows[0])]
    segments: List[ParsedSegment] = []
    chunk_rows: List[str] = []
    chunk_index = 0
    for row_index, row in enumerate(rows[1:], start=1):
        if not any(cell.strip() for cell in row):
            continue
        row_text = " | ".join(
            f"{header[column]}: {row[column].strip()}"
            for column in range(min(len(header), len(row)))
        )
        chunk_rows.append(row_text)
        if len(chunk_rows) >= 5:
            chunk_index += 1
            segments.append(
                ParsedSegment(
                    Path(file_name).stem,
                    f"table-row-block-{chunk_index}",
                    "\n".join([f"表头: {' | '.join(header)}", *chunk_rows]),
                )
            )
            chunk_rows = []
    if chunk_rows:
        chunk_index += 1
        segments.append(
            ParsedSegment(
                Path(file_name).stem,
                f"table-row-block-{chunk_index}",
                "\n".join([f"表头: {' | '.join(header)}", *chunk_rows]),
            )
        )
    return segments


def parse_plain_segments(file_name: str, text: str, source_uri: str) -> List[ParsedSegment]:
    paragraphs = [block.strip() for block in re.split(r"\n\s*\n", text) if block.strip()]
    if source_uri.startswith("http"):
        paragraphs = paragraphs[:200]
    segments: List[ParsedSegment] = []
    for paragraph in paragraphs:
        segments.extend(faq_or_plain_segments(Path(file_name).stem, "", paragraph))
    return segments


def faq_or_plain_segments(title: str, heading_path: str, text: str) -> List[ParsedSegment]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    segments: List[ParsedSegment] = []
    index = 0
    while index < len(lines):
        question_match = re.match(r"^(q|问)[:：]\s*(.+)$", lines[index], flags=re.IGNORECASE)
        if question_match and index + 1 < len(lines):
            answer_match = re.match(r"^(a|答)[:：]\s*(.+)$", lines[index + 1], flags=re.IGNORECASE)
            if answer_match:
                question = question_match.group(2).strip()
                answer = answer_match.group(2).strip()
                segments.append(ParsedSegment(question or title, heading_path, f"Q: {question}\nA: {answer}"))
                index += 2
                continue
        segments.extend(split_long_segment(ParsedSegment(title, heading_path, lines[index])))
        index += 1
    if segments:
        return segments
    return split_long_segment(ParsedSegment(title, heading_path, text.strip()))


def split_long_segment(segment: ParsedSegment) -> List[ParsedSegment]:
    if estimate_tokens(segment.content) <= 600:
        return [segment]
    sentences = split_sentences(segment.content)
    chunks: List[ParsedSegment] = []
    buffer: List[str] = []
    for sentence in sentences:
        candidate = "\n".join(buffer + [sentence]).strip()
        if buffer and estimate_tokens(candidate) > 400:
            chunks.append(
                ParsedSegment(
                    segment.title,
                    segment.heading_path,
                    "\n".join(buffer).strip(),
                    segment.page_number,
                )
            )
            overlap_sentences = split_sentences("\n".join(buffer))[-2:]
            buffer = overlap_sentences[:]
        buffer.append(sentence)
    if buffer:
        chunks.append(
            ParsedSegment(
                segment.title,
                segment.heading_path,
                "\n".join(buffer).strip(),
                segment.page_number,
            )
        )
    return chunks


def build_chunks(knowledge_base_id: str, document_id: str, title: str, source_uri: str, segments: List[ParsedSegment]) -> List[KnowledgeChunkRecord]:
    chunks: List[KnowledgeChunkRecord] = []
    for index, segment in enumerate(segments):
        if not segment.content.strip():
            continue
        chunks.append(
            KnowledgeChunkRecord(
                id=f"kb-chunk-{uuid.uuid4().hex[:12]}",
                knowledge_base_id=knowledge_base_id,
                document_id=document_id,
                chunk_index=index,
                title=segment.title or title,
                heading_path=segment.heading_path,
                source_uri=source_uri,
                page_number=segment.page_number,
                content=segment.content.strip(),
                token_count=estimate_tokens(segment.content),
                normalized_title=normalize_text(segment.title or title),
                normalized_heading_path=normalize_text(segment.heading_path),
                normalized_content=normalize_text(segment.content),
                normalized_source_uri=normalize_text(source_uri),
            )
        )
    return chunks


def trigram_similarity(left: str, right: str) -> float:
    if not left or not right:
        return 0.0

    def trigrams(value: str) -> set[str]:
        padded = f"  {value} "
        return {padded[index : index + 3] for index in range(len(padded) - 2)}

    left_trigrams = trigrams(left)
    right_trigrams = trigrams(right)
    if not left_trigrams or not right_trigrams:
        return 0.0
    overlap = len(left_trigrams & right_trigrams)
    return overlap / max(len(left_trigrams), len(right_trigrams))


def lexical_score(chunk: KnowledgeChunkRecord, query: str) -> float:
    query_terms = tokenize(query)
    if not query_terms:
        return 0.0
    title_hits = sum(1 for term in query_terms if term in chunk.normalized_title)
    heading_hits = sum(1 for term in query_terms if term in chunk.normalized_heading_path)
    body_hits = sum(1 for term in query_terms if term in chunk.normalized_content)
    metadata_hits = sum(1 for term in query_terms if term in chunk.normalized_source_uri)
    exact_phrase_boost = 1.5 if normalize_text(query) and normalize_text(query) in chunk.normalized_content else 0.0
    trigram = trigram_similarity(normalize_text(query), chunk.normalized_content)
    return round(
        (title_hits * 3.0)
        + (heading_hits * 2.0)
        + body_hits
        + (metadata_hits * 0.5)
        + (trigram * 4.0)
        + exact_phrase_boost,
        4,
    )


def term_weights(value: str) -> Counter[str]:
    tokens = tokenize(value)
    if not tokens:
        return Counter()
    counts = Counter(tokens)
    length = math.sqrt(sum(count * count for count in counts.values())) or 1.0
    return Counter({token: count / length for token, count in counts.items()})


def cosine_similarity(left: Counter[str], right: Counter[str]) -> float:
    if not left or not right:
        return 0.0
    return sum(left[token] * right.get(token, 0.0) for token in left)


def vector_score(chunk: KnowledgeChunkRecord, query: str) -> float:
    query_vector = term_weights(query)
    content_vector = term_weights(chunk.content)
    title_vector = term_weights(chunk.title)
    heading_vector = term_weights(chunk.heading_path)
    return round(
        (cosine_similarity(query_vector, title_vector) * 3.0)
        + (cosine_similarity(query_vector, heading_vector) * 2.0)
        + (cosine_similarity(query_vector, content_vector) * 4.0),
        4,
    )


def rerank_bonus(chunk: KnowledgeChunkRecord, query: str) -> float:
    query_terms = tokenize(query)
    if not query_terms:
        return 0.0
    matched_terms = [term for term in query_terms if term in chunk.normalized_content or term in chunk.normalized_title]
    if not matched_terms:
        return 0.0
    coverage = len(set(matched_terms)) / len(set(query_terms))
    title_exact = 0.8 if normalize_text(query) in chunk.normalized_title else 0.0
    heading_exact = 0.5 if normalize_text(query) in chunk.normalized_heading_path else 0.0
    density = min(1.0, len(matched_terms) / max(1, chunk.token_count / 30))
    return round((coverage * 1.8) + (density * 1.2) + title_exact + heading_exact, 4)


def final_score(chunk: KnowledgeChunkRecord, query: str, retrieval_mode: str) -> float:
    mode = normalize_retrieval_mode(retrieval_mode)
    lexical = lexical_score(chunk, query)
    vector = vector_score(chunk, query)
    if mode == "LEXICAL":
        return lexical
    if mode == "VECTOR":
        return vector
    hybrid = (lexical * 0.65) + (vector * 0.75)
    return round(hybrid + rerank_bonus(chunk, query), 4)


def snippet_for_query(content: str, query_terms: List[str]) -> str:
    normalized = normalize_text(content)
    index = min((normalized.find(term) for term in query_terms if term and normalized.find(term) >= 0), default=-1)
    if index < 0:
        return content[:220]
    start = max(0, index - 60)
    end = min(len(content), start + 220)
    return content[start:end]


def normalize_retrieval_mode(value: Optional[str]) -> str:
    normalized = (value or DEFAULT_SNAPSHOT_RETRIEVAL_MODE).strip().upper()
    if normalized not in SNAPSHOT_RETRIEVAL_MODES:
        return "HYBRID"
    return normalized


def infer_source_uri(file_record: KnowledgeFileRecord) -> str:
    if re.match(r"^https?://", file_record.file_name):
        return file_record.file_name
    return f"upload://{file_record.knowledge_base_id}/{file_record.file_name}"


def file_response(record: KnowledgeFileRecord) -> KnowledgeFileResponse:
    return KnowledgeFileResponse(
        id=record.id,
        knowledgeBaseId=record.knowledge_base_id,
        uploadSessionId=record.upload_session_id,
        fileName=record.file_name,
        contentType=record.content_type,
        sizeBytes=record.size_bytes,
        status=record.status,
        errorMessage=record.error_message,
        createdAt=record.created_at,
        updatedAt=record.updated_at,
    )


def import_job_response(record: KnowledgeImportJobRecord) -> ImportJobResponse:
    return ImportJobResponse(
        id=record.id,
        knowledgeBaseId=record.knowledge_base_id,
        fileId=record.file_id,
        status=record.status,
        failureReason=record.failure_reason,
        createdAt=record.created_at,
        updatedAt=record.updated_at,
        completedAt=record.completed_at,
    )


def document_response(db: Session, record: KnowledgeDocumentRecord) -> DocumentResponse:
    chunk_count = len(
        db.scalars(select(KnowledgeChunkRecord).where(KnowledgeChunkRecord.document_id == record.id)).all()
    )
    return DocumentResponse(
        id=record.id,
        knowledgeBaseId=record.knowledge_base_id,
        fileId=record.file_id,
        title=record.title,
        sourceUri=record.source_uri,
        documentType=record.document_type,
        status=record.status,
        chunkCount=chunk_count,
        createdAt=record.created_at,
        updatedAt=record.updated_at,
    )


def snapshot_response(record: IndexSnapshotRecord) -> IndexSnapshotResponse:
    return IndexSnapshotResponse(
        id=record.id,
        knowledgeBaseId=record.knowledge_base_id,
        retrievalBackend=record.retrieval_backend,
        retrievalMode=record.retrieval_mode,
        status=record.status,
        documentCount=record.document_count,
        chunkCount=record.chunk_count,
        failureReason=record.failure_reason,
        builtAt=record.built_at,
        createdAt=record.created_at,
        updatedAt=record.updated_at,
    )


def make_import_job(db: Session, knowledge_base_id: str, file_name: str, content_type: str, payload: bytes) -> dict:
    try:
        ensure_supported_document_type(file_name, content_type)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    upload_session = UploadSessionRecord(
        id=f"upload-session-{uuid.uuid4().hex[:10]}",
        knowledge_base_id=knowledge_base_id,
        status="COMPLETED",
    )
    object_key = f"{knowledge_base_id}/{uuid.uuid4().hex[:12]}-{sanitize_object_name(file_name)}"
    storage.put_bytes(object_key, payload, content_type)
    file_record = KnowledgeFileRecord(
        id=f"kb-file-{uuid.uuid4().hex[:12]}",
        knowledge_base_id=knowledge_base_id,
        upload_session_id=upload_session.id,
        file_name=file_name,
        content_type=content_type or "application/octet-stream",
        object_key=object_key,
        size_bytes=len(payload),
        status="UPLOADED",
    )
    import_job = KnowledgeImportJobRecord(
        id=f"kb-import-{uuid.uuid4().hex[:12]}",
        knowledge_base_id=knowledge_base_id,
        file_id=file_record.id,
        status="PENDING",
    )
    db.add_all([upload_session, file_record, import_job])
    db.commit()
    return {
        "file": file_response(file_record).model_dump(mode="json"),
        "importJob": import_job_response(import_job).model_dump(mode="json"),
    }


def sanitize_object_name(value: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9._-]+", "-", value).strip("-")
    return cleaned or "document.bin"


def infer_url_file_name(url: str, title: Optional[str], content_type: str) -> str:
    parsed = urlparse(url)
    leaf = Path(parsed.path).name or parsed.netloc or "page"
    leaf = leaf.split("?")[0]
    base = title.strip() if title and title.strip() else leaf
    base = sanitize_object_name(base)
    suffix = Path(base).suffix.lower()
    if suffix in {".md", ".txt", ".html", ".csv", ".pdf", ".docx"}:
        return base
    if "html" in content_type.lower():
        return f"{base}.html"
    if "markdown" in content_type.lower():
        return f"{base}.md"
    if "csv" in content_type.lower():
        return f"{base}.csv"
    if "pdf" in content_type.lower():
        return f"{base}.pdf"
    return f"{base}.txt"


def fetch_url_payload(url: str) -> tuple[bytes, str]:
    request = Request(url, headers={"User-Agent": URL_IMPORT_USER_AGENT})
    try:
        with urlopen(request, timeout=URL_IMPORT_TIMEOUT_SECONDS) as response:
            content_type = response.headers.get("Content-Type", "text/html")
            payload = response.read()
            return payload, content_type
    except HTTPError as exc:
        raise HTTPException(status_code=400, detail=f"url import failed with http {exc.code}: {url}") from exc
    except URLError as exc:
        raise HTTPException(status_code=400, detail=f"url import failed: {exc.reason}") from exc


@app.on_event("startup")
def startup() -> None:
    Base.metadata.create_all(bind=engine)


@app.post("/internal/upload-sessions", response_model=UploadSessionResponse)
def create_upload_session(request: CreateUploadSessionRequest, db: Session = Depends(get_db)) -> UploadSessionResponse:
    session_record = UploadSessionRecord(
        id=f"upload-session-{uuid.uuid4().hex[:10]}",
        knowledge_base_id=request.knowledgeBaseId,
        status="OPEN",
    )
    db.add(session_record)
    db.commit()
    return UploadSessionResponse(
        id=session_record.id,
        knowledgeBaseId=session_record.knowledge_base_id,
        status=session_record.status,
        acceptedTypes=SUPPORTED_FILE_TYPES,
    )


@app.post("/internal/uploads")
def complete_upload(request: CompleteUploadRequest, db: Session = Depends(get_db)) -> dict:
    upload_session = db.get(UploadSessionRecord, request.uploadSessionId)
    if upload_session is None or upload_session.knowledge_base_id != request.knowledgeBaseId:
        raise HTTPException(status_code=404, detail="upload session not found")

    payload = base64.b64decode(request.contentBase64)
    try:
        ensure_supported_document_type(request.fileName, request.contentType)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    upload_session.status = "COMPLETED"
    upload_session.updated_at = now_utc()

    object_key = f"{request.knowledgeBaseId}/{uuid.uuid4().hex[:12]}-{sanitize_object_name(request.fileName)}"
    storage.put_bytes(object_key, payload, request.contentType)
    file_record = KnowledgeFileRecord(
        id=f"kb-file-{uuid.uuid4().hex[:12]}",
        knowledge_base_id=request.knowledgeBaseId,
        upload_session_id=request.uploadSessionId,
        file_name=request.fileName,
        content_type=request.contentType or "application/octet-stream",
        object_key=object_key,
        size_bytes=len(payload),
        status="UPLOADED",
    )
    import_job = KnowledgeImportJobRecord(
        id=f"kb-import-{uuid.uuid4().hex[:12]}",
        knowledge_base_id=request.knowledgeBaseId,
        file_id=file_record.id,
        status="PENDING",
    )
    db.add_all([file_record, import_job])
    db.commit()
    return {
        "file": file_response(file_record).model_dump(mode="json"),
        "importJob": import_job_response(import_job).model_dump(mode="json"),
    }


@app.post("/internal/url-imports")
def create_url_import(request: CreateUrlImportRequest, db: Session = Depends(get_db)) -> dict:
    payload, content_type = fetch_url_payload(request.url)
    file_name = infer_url_file_name(request.url, request.title, content_type)
    try:
        ensure_supported_document_type(file_name, content_type)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return make_import_job(db, request.knowledgeBaseId, request.url if request.url else file_name, content_type, payload)


@app.get("/internal/knowledge-bases/{knowledge_base_id}/files", response_model=List[KnowledgeFileResponse])
def list_files(knowledge_base_id: str, db: Session = Depends(get_db)) -> List[KnowledgeFileResponse]:
    records = db.scalars(
        select(KnowledgeFileRecord)
        .where(KnowledgeFileRecord.knowledge_base_id == knowledge_base_id)
        .order_by(KnowledgeFileRecord.created_at.desc())
    ).all()
    return [file_response(record) for record in records]


@app.get("/internal/knowledge-bases/{knowledge_base_id}/import-jobs", response_model=List[ImportJobResponse])
def list_import_jobs(knowledge_base_id: str, db: Session = Depends(get_db)) -> List[ImportJobResponse]:
    records = db.scalars(
        select(KnowledgeImportJobRecord)
        .where(KnowledgeImportJobRecord.knowledge_base_id == knowledge_base_id)
        .order_by(KnowledgeImportJobRecord.created_at.desc())
    ).all()
    return [import_job_response(record) for record in records]


@app.get("/internal/knowledge-bases/{knowledge_base_id}/documents", response_model=List[DocumentResponse])
def list_documents(knowledge_base_id: str, db: Session = Depends(get_db)) -> List[DocumentResponse]:
    records = db.scalars(
        select(KnowledgeDocumentRecord)
        .where(KnowledgeDocumentRecord.knowledge_base_id == knowledge_base_id)
        .order_by(KnowledgeDocumentRecord.created_at.desc())
    ).all()
    return [document_response(db, record) for record in records]


@app.post("/internal/knowledge-bases/{knowledge_base_id}/index-snapshots", response_model=IndexSnapshotResponse)
def create_index_snapshot(knowledge_base_id: str, request: CreateIndexSnapshotRequest, db: Session = Depends(get_db)) -> IndexSnapshotResponse:
    snapshot = IndexSnapshotRecord(
        id=f"snapshot-{uuid.uuid4().hex[:12]}",
        knowledge_base_id=knowledge_base_id,
        retrieval_backend="OPENSEARCH",
        retrieval_mode=normalize_retrieval_mode(request.retrievalMode),
        status="PENDING",
    )
    db.add(snapshot)
    db.flush()
    document_ids = list(dict.fromkeys(request.documentIds or []))
    if document_ids:
        db.add_all(
            [
                SnapshotDocumentSelectionRecord(
                    id=f"snapshot-document-{uuid.uuid4().hex[:12]}",
                    snapshot_id=snapshot.id,
                    document_id=document_id,
                )
                for document_id in document_ids
            ]
        )
    db.commit()
    return snapshot_response(snapshot)


@app.get("/internal/knowledge-bases/{knowledge_base_id}/index-snapshots", response_model=List[IndexSnapshotResponse])
def list_index_snapshots(knowledge_base_id: str, db: Session = Depends(get_db)) -> List[IndexSnapshotResponse]:
    records = db.scalars(
        select(IndexSnapshotRecord)
        .where(IndexSnapshotRecord.knowledge_base_id == knowledge_base_id)
        .order_by(IndexSnapshotRecord.created_at.desc())
    ).all()
    return [snapshot_response(record) for record in records]


@app.post("/internal/import-jobs/{job_id}/run", response_model=ImportJobResponse)
def run_import_job(job_id: str, db: Session = Depends(get_db)) -> ImportJobResponse:
    import_job = db.get(KnowledgeImportJobRecord, job_id)
    if import_job is None:
        raise HTTPException(status_code=404, detail="import job not found")
    file_record = db.get(KnowledgeFileRecord, import_job.file_id)
    if file_record is None:
        raise HTTPException(status_code=404, detail="knowledge file not found")

    import_job.status = "RUNNING"
    import_job.updated_at = now_utc()
    file_record.status = "IMPORTING"
    file_record.updated_at = now_utc()
    db.commit()

    try:
        payload = storage.get_bytes(file_record.object_key)
        parsed = parse_document(file_record, payload)
        if not parsed.body_text.strip():
            raise ValueError("no extractable text found")

        existing_documents = db.scalars(
            select(KnowledgeDocumentRecord).where(KnowledgeDocumentRecord.file_id == file_record.id)
        ).all()
        for existing_document in existing_documents:
            db.query(KnowledgeChunkRecord).filter(KnowledgeChunkRecord.document_id == existing_document.id).delete()
            db.delete(existing_document)
        db.flush()

        document = KnowledgeDocumentRecord(
            id=f"kb-document-{uuid.uuid4().hex[:12]}",
            knowledge_base_id=file_record.knowledge_base_id,
            file_id=file_record.id,
            title=parsed.title,
            source_uri=infer_source_uri(file_record),
            document_type=parsed.document_type,
            status="READY",
            body_text=parsed.body_text,
        )
        chunks = build_chunks(file_record.knowledge_base_id, document.id, document.title, document.source_uri, parsed.segments)
        if not chunks:
            raise ValueError("no chunks generated from document")

        db.add(document)
        db.add_all(chunks)
        import_job.status = "COMPLETED"
        import_job.completed_at = now_utc()
        import_job.failure_reason = None
        file_record.status = "IMPORTED"
        file_record.error_message = None
    except Exception as exc:
        import_job.status = "FAILED"
        import_job.failure_reason = str(exc)
        file_record.status = "FAILED"
        file_record.error_message = str(exc)
    finally:
        import_job.updated_at = now_utc()
        file_record.updated_at = now_utc()
        db.commit()
    return import_job_response(import_job)


@app.post("/internal/index-snapshots/{snapshot_id}/build", response_model=IndexSnapshotResponse)
def build_index_snapshot(snapshot_id: str, db: Session = Depends(get_db)) -> IndexSnapshotResponse:
    snapshot = db.get(IndexSnapshotRecord, snapshot_id)
    if snapshot is None:
        raise HTTPException(status_code=404, detail="index snapshot not found")
    if not opensearch.enabled:
        raise HTTPException(status_code=500, detail="opensearch is not configured for knowledge snapshot builds")

    snapshot.status = "BUILDING"
    snapshot.updated_at = now_utc()
    db.query(IndexSnapshotChunkRecord).filter(IndexSnapshotChunkRecord.snapshot_id == snapshot_id).delete()
    db.commit()

    try:
        selected_document_ids = [
            selection.document_id
            for selection in db.scalars(
                select(SnapshotDocumentSelectionRecord).where(SnapshotDocumentSelectionRecord.snapshot_id == snapshot.id)
            ).all()
        ]
        query = select(KnowledgeDocumentRecord).where(
            KnowledgeDocumentRecord.knowledge_base_id == snapshot.knowledge_base_id,
            KnowledgeDocumentRecord.status == "READY",
        )
        documents = db.scalars(query).all()
        if selected_document_ids:
            documents = [document for document in documents if document.id in set(selected_document_ids)]
        if not documents:
            raise ValueError("no ready documents available for snapshot")

        all_chunks: List[KnowledgeChunkRecord] = []
        chunk_records: List[IndexSnapshotChunkRecord] = []
        chunk_count = 0
        for document in documents:
            chunks = db.scalars(
                select(KnowledgeChunkRecord)
                .where(KnowledgeChunkRecord.document_id == document.id)
                .order_by(KnowledgeChunkRecord.chunk_index.asc())
            ).all()
            for chunk in chunks:
                all_chunks.append(chunk)
                chunk_records.append(
                    IndexSnapshotChunkRecord(
                        id=f"snapshot-chunk-{uuid.uuid4().hex[:12]}",
                        snapshot_id=snapshot.id,
                        chunk_id=chunk.id,
                        document_id=document.id,
                    )
                )
            chunk_count += len(chunks)

        if chunk_count == 0:
            raise ValueError("no chunks available for snapshot")

        opensearch.bulk_index_chunks(snapshot.id, snapshot.knowledge_base_id, all_chunks)
        db.add_all(chunk_records)
        snapshot.status = "READY"
        snapshot.document_count = len(documents)
        snapshot.chunk_count = chunk_count
        snapshot.retrieval_backend = "OPENSEARCH"
        snapshot.failure_reason = None
        snapshot.built_at = now_utc()
    except Exception as exc:
        snapshot.status = "FAILED"
        snapshot.failure_reason = str(exc)
        snapshot.document_count = 0
        snapshot.chunk_count = 0
    finally:
        snapshot.updated_at = now_utc()
        db.commit()
    return snapshot_response(snapshot)


@app.get("/internal/index-snapshots/{snapshot_id}", response_model=IndexSnapshotResponse)
def get_index_snapshot(snapshot_id: str, db: Session = Depends(get_db)) -> IndexSnapshotResponse:
    snapshot = db.get(IndexSnapshotRecord, snapshot_id)
    if snapshot is None:
        raise HTTPException(status_code=404, detail="index snapshot not found")
    return snapshot_response(snapshot)


@app.post("/internal/retrieve", response_model=RetrieveResponse)
def retrieve(request: RetrieveRequest, db: Session = Depends(get_db)) -> RetrieveResponse:
    snapshot = db.get(IndexSnapshotRecord, request.indexSnapshotId)
    if snapshot is None or snapshot.status != "READY":
        raise HTTPException(status_code=404, detail="index snapshot is not ready")
    if snapshot.retrieval_backend != "OPENSEARCH":
        raise HTTPException(status_code=500, detail=f"unsupported retrieval backend: {snapshot.retrieval_backend}")
    if not opensearch.enabled:
        raise HTTPException(status_code=500, detail="opensearch is not configured for retrieval")

    retrieval_mode = normalize_retrieval_mode(request.retrievalMode or snapshot.retrieval_mode)
    candidate_ids = opensearch.search_chunk_ids(
        snapshot.id,
        request.query,
        retrieval_mode,
        max(20, min(MAX_RERANK_CANDIDATES, request.topK * 8)),
    )
    if not candidate_ids:
        return RetrieveResponse(hits=[], lowConfidence=True)

    mappings = db.scalars(select(IndexSnapshotChunkRecord).where(IndexSnapshotChunkRecord.snapshot_id == snapshot.id)).all()
    mapping_index = {mapping.chunk_id: mapping for mapping in mappings}
    candidates: List[RetrieveHit] = []
    for chunk_id in candidate_ids:
        mapping = mapping_index.get(chunk_id)
        if mapping is None:
            continue
        chunk = db.get(KnowledgeChunkRecord, mapping.chunk_id)
        document = db.get(KnowledgeDocumentRecord, mapping.document_id)
        if chunk is None or document is None:
            continue
        score = final_score(chunk, request.query, retrieval_mode)
        if score < request.minScore:
            continue
        candidates.append(
            RetrieveHit(
                chunkId=chunk.id,
                documentId=document.id,
                documentTitle=document.title,
                sourceUri=chunk.source_uri,
                snippet=snippet_for_query(chunk.content, tokenize(request.query)),
                score=round(score, 4),
                pageNumber=chunk.page_number,
                headingPath=chunk.heading_path,
            )
        )
    candidates.sort(key=lambda item: item.score, reverse=True)
    hits = candidates[: max(1, min(request.topK, MAX_RERANK_CANDIDATES))]
    return RetrieveResponse(hits=hits, lowConfidence=not hits)
