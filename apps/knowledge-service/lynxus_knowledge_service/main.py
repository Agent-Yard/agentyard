from __future__ import annotations

import base64
import csv
import io
import json
import logging
import math
import os
import re
import secrets
import uuid
from contextlib import asynccontextmanager
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, AsyncIterator, Generator, List, Optional, Sequence, Union
from urllib.error import HTTPError, URLError
from urllib.parse import urlparse
from urllib.request import Request as UrlRequest, urlopen

from fastapi import Depends, FastAPI, Header, HTTPException, Request as FastAPIRequest
from fastapi.responses import JSONResponse
from lynxus_knowledge_service.object_storage import Storage, load_storage_settings
from lynxus_common import (
    ReadinessCheck,
    TRACEPARENT_HEADER,
    bind_request_log_context,
    build_readiness_report,
    clear_log_context,
    configure_structured_logging,
)
from pydantic import BaseModel, Field
from sqlalchemy import Boolean, DateTime, ForeignKey, Integer, String, Text, bindparam, cast, create_engine, func, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column, relationship, sessionmaker
from sqlalchemy.sql.type_api import UserDefinedType

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover
    DocxDocument = None

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover
    PdfReader = None

def now_utc() -> datetime:
    return datetime.now(timezone.utc)


SERVICE_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DATA_ROOT = SERVICE_ROOT / "data"
DEFAULT_STORAGE_ROOT = DEFAULT_DATA_ROOT / "storage"
DEFAULT_DATABASE_URL = "postgresql+psycopg://lynxus:lynxus@127.0.0.1:5432/lynxus_knowledge"


DATABASE_URL = os.getenv("LYNXUS_KNOWLEDGE_DATABASE_URL", DEFAULT_DATABASE_URL)
STORAGE_SETTINGS = load_storage_settings(DEFAULT_STORAGE_ROOT)
URL_IMPORT_TIMEOUT_SECONDS = float(os.getenv("LYNXUS_KNOWLEDGE_URL_IMPORT_TIMEOUT_SECONDS", "15"))
URL_IMPORT_USER_AGENT = os.getenv(
    "LYNXUS_KNOWLEDGE_URL_IMPORT_USER_AGENT",
    "LynxusKnowledgeService/2.0",
)
EMBEDDING_BASE_URL = os.getenv("LYNXUS_KNOWLEDGE_EMBEDDING_BASE_URL", "").rstrip("/")
EMBEDDING_MODEL = os.getenv("LYNXUS_KNOWLEDGE_EMBEDDING_MODEL", "").strip()
EMBEDDING_API_KEY = os.getenv("LYNXUS_KNOWLEDGE_EMBEDDING_API_KEY", "").strip()
EMBEDDING_DIMENSIONS = int(os.getenv("LYNXUS_KNOWLEDGE_EMBEDDING_DIMENSIONS", "1024"))
EMBEDDING_BATCH_SIZE = max(1, int(os.getenv("LYNXUS_KNOWLEDGE_EMBEDDING_BATCH_SIZE", "16")))
EMBEDDING_TIMEOUT_SECONDS = float(os.getenv("LYNXUS_KNOWLEDGE_EMBEDDING_TIMEOUT_SECONDS", "15"))
DEFAULT_SNAPSHOT_RETRIEVAL_MODE = os.getenv("LYNXUS_KNOWLEDGE_DEFAULT_RETRIEVAL_MODE", "HYBRID").upper()
DEFAULT_SNAPSHOT_RETRIEVAL_BACKEND = "PGVECTOR"
INSTANCE_ID = (os.getenv("LYNXUS_INSTANCE_ID") or "lynxus-knowledge-service").strip() or "lynxus-knowledge-service"

engine = create_engine(DATABASE_URL, future=True)
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, expire_on_commit=False)
configure_structured_logging("lynxus-knowledge-service", "LYNXUS_KNOWLEDGE_SERVICE_LOG_LEVEL")
logger = logging.getLogger(__name__)

SUPPORTED_FILE_TYPES = ["pdf", "docx", "md", "txt", "html", "csv"]
SNAPSHOT_RETRIEVAL_MODES = {"LEXICAL", "VECTOR", "HYBRID"}
MAX_RERANK_CANDIDATES = 20
LEXICAL_CANDIDATE_MULTIPLIER = 8


class TSVectorType(UserDefinedType):
    cache_ok = True

    def get_col_spec(self, **_: Any) -> str:
        return "tsvector"


class VectorType(UserDefinedType):
    cache_ok = True

    def __init__(self, dimensions: int) -> None:
        self.dimensions = dimensions

    def get_col_spec(self, **_: Any) -> str:
        return f"vector({self.dimensions})"

    def bind_processor(self, dialect):
        def process(value: Optional[Sequence[float]]) -> Optional[str]:
            if value is None:
                return None
            return "[" + ",".join(f"{float(item):.12f}" for item in value) + "]"

        return process

    def bind_expression(self, bindvalue):
        return cast(bindvalue, self)

    def result_processor(self, dialect, coltype):
        def process(value: Any) -> Optional[List[float]]:
            if value is None:
                return None
            if isinstance(value, str):
                stripped = value.strip()[1:-1].strip()
                if not stripped:
                    return []
                return [float(item) for item in stripped.split(",")]
            if isinstance(value, (list, tuple)):
                return [float(item) for item in value]
            return None

        return process


def internal_auth_token() -> str:
    token = os.getenv("LYNXUS_INTERNAL_AUTH_TOKEN", "").strip()
    if not token:
        raise RuntimeError("LYNXUS_INTERNAL_AUTH_TOKEN must be configured")
    return token


def require_internal_bearer(authorization: str | None = Header(default=None)) -> None:
    expected = internal_auth_token()
    if authorization is None or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="internal authentication is required")
    actual = authorization.removeprefix("Bearer ").strip()
    if not actual or not secrets.compare_digest(actual, expected):
        raise HTTPException(status_code=401, detail="invalid internal authentication token")


class Base(DeclarativeBase):
    pass


class UploadSessionRecord(Base):
    __tablename__ = "knowledge_upload_session"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    knowledge_base_id: Mapped[str] = mapped_column(String(64), index=True)
    status: Mapped[str] = mapped_column(String(32), default="OPEN")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)


class KnowledgeFileRecord(Base):
    __tablename__ = "knowledge_file"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    knowledge_base_id: Mapped[str] = mapped_column(String(64), index=True)
    upload_session_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_upload_session.id"))
    source_type: Mapped[str] = mapped_column(String(32), default="FILE_UPLOAD")
    source_uri: Mapped[str] = mapped_column(String(1024), default="")
    file_name: Mapped[str] = mapped_column(String(255))
    content_type: Mapped[str] = mapped_column(String(128))
    object_key: Mapped[str] = mapped_column(String(255), default="")
    size_bytes: Mapped[int] = mapped_column(Integer, default=0)
    status: Mapped[str] = mapped_column(String(32), default="UPLOADED")
    error_message: Mapped[Optional[str]] = mapped_column(Text, nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)

    import_jobs: Mapped[List["KnowledgeImportJobRecord"]] = relationship(back_populates="knowledge_file")


class KnowledgeImportJobRecord(Base):
    __tablename__ = "knowledge_import_job"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    knowledge_base_id: Mapped[str] = mapped_column(String(64), index=True)
    file_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_file.id"))
    status: Mapped[str] = mapped_column(String(32), default="QUEUED")
    stage: Mapped[str] = mapped_column(String(64), default="QUEUED")
    progress_percent: Mapped[int] = mapped_column(Integer, default=0)
    retry_count: Mapped[int] = mapped_column(Integer, default=0)
    retryable: Mapped[bool] = mapped_column(Boolean, default=False)
    failure_reason: Mapped[Optional[str]] = mapped_column(Text, nullable=True)
    started_at: Mapped[Optional[datetime]] = mapped_column(DateTime(timezone=True), nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    completed_at: Mapped[Optional[datetime]] = mapped_column(DateTime(timezone=True), nullable=True)

    knowledge_file: Mapped[KnowledgeFileRecord] = relationship(back_populates="import_jobs")


class KnowledgeDocumentRecord(Base):
    __tablename__ = "knowledge_document"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    knowledge_base_id: Mapped[str] = mapped_column(String(64), index=True)
    file_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_file.id"))
    title: Mapped[str] = mapped_column(String(255))
    source_uri: Mapped[str] = mapped_column(String(255))
    document_type: Mapped[str] = mapped_column(String(32))
    status: Mapped[str] = mapped_column(String(32), default="READY")
    body_text: Mapped[str] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)

    chunks: Mapped[List["KnowledgeChunkRecord"]] = relationship(back_populates="document")


class KnowledgeChunkRecord(Base):
    __tablename__ = "knowledge_chunk"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    knowledge_base_id: Mapped[str] = mapped_column(String(64), index=True)
    document_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_document.id"))
    chunk_index: Mapped[int] = mapped_column(Integer)
    title: Mapped[str] = mapped_column(String(255))
    heading_path: Mapped[str] = mapped_column(String(512), default="")
    source_uri: Mapped[str] = mapped_column(String(255), default="")
    page_number: Mapped[Optional[int]] = mapped_column(Integer, nullable=True)
    content: Mapped[str] = mapped_column(Text)
    token_count: Mapped[int] = mapped_column(Integer)
    normalized_title: Mapped[str] = mapped_column(Text)
    normalized_heading_path: Mapped[str] = mapped_column(Text)
    normalized_content: Mapped[str] = mapped_column(Text)
    normalized_source_uri: Mapped[str] = mapped_column(Text)
    search_text: Mapped[str] = mapped_column(Text)
    search_vector: Mapped[Optional[str]] = mapped_column(TSVectorType(), nullable=True)
    embedding: Mapped[Optional[List[float]]] = mapped_column(VectorType(EMBEDDING_DIMENSIONS), nullable=True)
    embedding_model: Mapped[Optional[str]] = mapped_column(String(255), nullable=True)
    embedding_dimensions: Mapped[Optional[int]] = mapped_column(Integer, nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)

    document: Mapped[KnowledgeDocumentRecord] = relationship(back_populates="chunks")


class IndexSnapshotRecord(Base):
    __tablename__ = "knowledge_index_snapshot"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    knowledge_base_id: Mapped[str] = mapped_column(String(64), index=True)
    retrieval_backend: Mapped[str] = mapped_column(String(64), default=DEFAULT_SNAPSHOT_RETRIEVAL_BACKEND)
    retrieval_mode: Mapped[str] = mapped_column(String(32), default=DEFAULT_SNAPSHOT_RETRIEVAL_MODE)
    status: Mapped[str] = mapped_column(String(32), default="QUEUED")
    stage: Mapped[str] = mapped_column(String(64), default="QUEUED")
    progress_percent: Mapped[int] = mapped_column(Integer, default=0)
    retry_count: Mapped[int] = mapped_column(Integer, default=0)
    retryable: Mapped[bool] = mapped_column(Boolean, default=False)
    document_count: Mapped[int] = mapped_column(Integer, default=0)
    chunk_count: Mapped[int] = mapped_column(Integer, default=0)
    failure_reason: Mapped[Optional[str]] = mapped_column(Text, nullable=True)
    started_at: Mapped[Optional[datetime]] = mapped_column(DateTime(timezone=True), nullable=True)
    built_at: Mapped[Optional[datetime]] = mapped_column(DateTime(timezone=True), nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)

    snapshot_chunks: Mapped[List["IndexSnapshotChunkRecord"]] = relationship(back_populates="snapshot")
    selected_documents: Mapped[List["SnapshotDocumentSelectionRecord"]] = relationship(back_populates="snapshot")


class IndexSnapshotChunkRecord(Base):
    __tablename__ = "knowledge_index_snapshot_chunk"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    snapshot_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_index_snapshot.id"), index=True)
    chunk_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_chunk.id"))
    document_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_document.id"))

    snapshot: Mapped[IndexSnapshotRecord] = relationship(back_populates="snapshot_chunks")


class SnapshotDocumentSelectionRecord(Base):
    __tablename__ = "knowledge_snapshot_document_selection"

    id: Mapped[str] = mapped_column(String(64), primary_key=True)
    snapshot_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_index_snapshot.id"), index=True)
    document_id: Mapped[str] = mapped_column(String(64), ForeignKey("knowledge_document.id"))

    snapshot: Mapped[IndexSnapshotRecord] = relationship(back_populates="selected_documents")


class CreateUploadSessionRequest(BaseModel):
    knowledgeBaseId: str


class UploadSessionResponse(BaseModel):
    id: str
    knowledgeBaseId: str
    status: str
    acceptedTypes: List[str]


class CompleteUploadRequest(BaseModel):
    knowledgeBaseId: str
    uploadSessionId: str
    fileName: str
    contentType: str
    contentBase64: str


class CreateUrlImportRequest(BaseModel):
    knowledgeBaseId: str
    url: str
    title: Optional[str] = None


class KnowledgeFileResponse(BaseModel):
    id: str
    knowledgeBaseId: str
    uploadSessionId: str
    sourceType: str
    sourceUri: str
    fileName: str
    contentType: str
    sizeBytes: int
    status: str
    errorMessage: Optional[str] = None
    createdAt: datetime
    updatedAt: datetime


class ImportJobResponse(BaseModel):
    id: str
    knowledgeBaseId: str
    fileId: str
    sourceType: str
    sourceUri: str
    fileName: str
    status: str
    stage: str
    progressPercent: int
    retryCount: int
    retryable: bool
    failureReason: Optional[str] = None
    startedAt: Optional[datetime] = None
    createdAt: datetime
    updatedAt: datetime
    completedAt: Optional[datetime] = None


class DocumentResponse(BaseModel):
    id: str
    knowledgeBaseId: str
    fileId: str
    title: str
    sourceUri: str
    documentType: str
    status: str
    chunkCount: int
    createdAt: datetime
    updatedAt: datetime


class DocumentDeletionBlockerResponse(BaseModel):
    snapshotId: str
    status: str
    stage: str
    retrievalMode: str
    reason: str


class DocumentDeletionPreviewResponse(BaseModel):
    documentId: str
    knowledgeBaseId: str
    fileId: str
    fileName: str
    sourceUri: str
    title: str
    chunkCount: int
    canDelete: bool
    blockers: List[DocumentDeletionBlockerResponse]


class DocumentDeletionResponse(BaseModel):
    documentId: str
    knowledgeBaseId: str
    fileId: str
    fileName: str
    title: str
    deletedChunkCount: int
    deletedImportJobCount: int
    deletedDocumentCount: int
    deletedStorageObject: bool


class CreateIndexSnapshotRequest(BaseModel):
    knowledgeBaseId: Optional[str] = None
    documentIds: List[str] = Field(default_factory=list)
    retrievalMode: str = Field(default=DEFAULT_SNAPSHOT_RETRIEVAL_MODE)


class IndexSnapshotResponse(BaseModel):
    id: str
    knowledgeBaseId: str
    retrievalBackend: str
    retrievalMode: str
    status: str
    stage: str
    progressPercent: int
    retryCount: int
    retryable: bool
    documentCount: int
    chunkCount: int
    failureReason: Optional[str] = None
    startedAt: Optional[datetime] = None
    builtAt: Optional[datetime] = None
    createdAt: datetime
    updatedAt: datetime


class RetrieveRequest(BaseModel):
    indexSnapshotId: str
    query: str
    topK: int = 5
    minScore: float = 0.1
    retrievalMode: Optional[str] = None


class RetrieveHit(BaseModel):
    chunkId: str
    documentId: str
    documentTitle: str
    sourceUri: str
    snippet: str
    score: float
    pageNumber: Optional[int] = None
    headingPath: str = ""


class RetrieveResponse(BaseModel):
    hits: List[RetrieveHit]
    lowConfidence: bool


class ReadChunksRequest(BaseModel):
    indexSnapshotId: str
    chunkIds: List[str] = Field(default_factory=list)


class ReadChunk(BaseModel):
    chunkId: str
    documentId: str
    documentTitle: str
    sourceUri: str
    headingPath: str = ""
    pageNumber: Optional[int] = None
    content: str


class ReadChunksResponse(BaseModel):
    chunks: List[ReadChunk]


@dataclass
class ParsedDocument:
    title: str
    document_type: str
    body_text: str
    segments: List["ParsedSegment"]


@dataclass
class ParsedSegment:
    title: str
    heading_path: str
    content: str
    page_number: Optional[int] = None


class EmbeddingClient:
    def __init__(self) -> None:
        self.base_url = EMBEDDING_BASE_URL
        self.model = EMBEDDING_MODEL
        self.api_key = EMBEDDING_API_KEY
        self.dimensions = EMBEDDING_DIMENSIONS
        self.batch_size = EMBEDDING_BATCH_SIZE
        self.timeout_seconds = EMBEDDING_TIMEOUT_SECONDS

    def validate_configuration(self) -> None:
        missing = [
            name
            for name, value in {
                "LYNXUS_KNOWLEDGE_EMBEDDING_BASE_URL": self.base_url,
                "LYNXUS_KNOWLEDGE_EMBEDDING_MODEL": self.model,
                "LYNXUS_KNOWLEDGE_EMBEDDING_API_KEY": self.api_key,
            }.items()
            if not value
        ]
        if self.dimensions <= 0:
            missing.append("LYNXUS_KNOWLEDGE_EMBEDDING_DIMENSIONS")
        if missing:
            raise RuntimeError("embedding provider is not configured: " + ", ".join(missing))

    def request(self, payload: dict) -> dict:
        self.validate_configuration()
        headers = {
            "Accept": "application/json",
            "Content-Type": "application/json",
            "Authorization": f"Bearer {self.api_key}",
        }
        request = UrlRequest(
            f"{self.base_url}/embeddings",
            data=json.dumps(payload).encode("utf-8"),
            method="POST",
            headers=headers,
        )
        try:
            with urlopen(request, timeout=self.timeout_seconds) as response:
                body = response.read().decode("utf-8", errors="ignore").strip()
                return json.loads(body) if body else {}
        except HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="ignore")
            raise ValueError(f"embedding request failed: {exc.code} {detail}") from exc
        except URLError as exc:
            raise ValueError(f"embedding request failed: {exc.reason}") from exc

    def embed_texts(self, texts: Sequence[str]) -> List[List[float]]:
        self.validate_configuration()
        if not texts:
            return []
        embeddings: List[List[float]] = []
        for start in range(0, len(texts), self.batch_size):
            batch = list(texts[start : start + self.batch_size])
            payload = self.request(
                {
                    "model": self.model,
                    "input": batch,
                    "dimensions": self.dimensions,
                    "encoding_format": "float",
                }
            )
            batch_embeddings = payload.get("data", [])
            if len(batch_embeddings) != len(batch):
                raise ValueError("embedding provider returned mismatched batch size")
            for item in batch_embeddings:
                vector = item.get("embedding")
                if not isinstance(vector, list) or len(vector) != self.dimensions:
                    raise ValueError("embedding provider returned an invalid embedding vector")
                embeddings.append([float(value) for value in vector])
        return embeddings


class PostgresRetrievalStore:
    def __init__(self, dimensions: int, embedding_model: str) -> None:
        self.dimensions = dimensions
        self.embedding_model = embedding_model

    @staticmethod
    def _normalize_scores(raw_scores: dict[str, float]) -> dict[str, float]:
        if not raw_scores:
            return {}
        values = list(raw_scores.values())
        min_score = min(values)
        max_score = max(values)
        if math.isclose(max_score, min_score):
            return {key: (1.0 if value > 0 else 0.0) for key, value in raw_scores.items()}
        return {
            key: max(0.0, min(1.0, (value - min_score) / (max_score - min_score)))
            for key, value in raw_scores.items()
        }

    @staticmethod
    def _vector_literal(embedding: Sequence[float]) -> str:
        return "[" + ",".join(f"{float(value):.12f}" for value in embedding) + "]"

    def refresh_chunk_search_vectors(self, db: Session, chunk_ids: Sequence[str]) -> None:
        if not chunk_ids:
            return
        statement = text(
            """
            update knowledge_chunk
               set search_vector =
                   setweight(to_tsvector('simple', coalesce(title, '')), 'A')
                || setweight(to_tsvector('simple', coalesce(heading_path, '')), 'B')
                || setweight(to_tsvector('simple', coalesce(content, '')), 'C')
                || setweight(to_tsvector('simple', coalesce(source_uri, '')), 'D')
             where id in :chunk_ids
            """
        ).bindparams(bindparam("chunk_ids", expanding=True))
        db.execute(statement, {"chunk_ids": list(chunk_ids)})

    def ensure_snapshot_chunks_compatible(self, db: Session, snapshot_id: str) -> None:
        row = db.execute(
            text(
                """
                select count(*) as invalid_count
                  from knowledge_index_snapshot_chunk mapping
                  join knowledge_chunk chunk on chunk.id = mapping.chunk_id
                 where mapping.snapshot_id = :snapshot_id
                   and (
                        chunk.embedding is null
                     or chunk.embedding_model is distinct from :embedding_model
                     or chunk.embedding_dimensions is distinct from :embedding_dimensions
                   )
                """
            ),
            {
                "snapshot_id": snapshot_id,
                "embedding_model": self.embedding_model,
                "embedding_dimensions": self.dimensions,
            },
        ).mappings().one()
        if int(row["invalid_count"]) > 0:
            raise ValueError("snapshot chunks were embedded with a different embedding configuration; re-import the documents")

    def lexical_candidates(self, db: Session, snapshot_id: str, query: str, size: int) -> dict[str, float]:
        normalized_query = normalize_text(query)
        rows = db.execute(
            text(
                """
                select chunk.id as chunk_id,
                       (
                           ts_rank_cd(
                               chunk.search_vector,
                               websearch_to_tsquery('simple', :query)
                           ) * 0.7
                         + greatest(
                               similarity(chunk.normalized_title, :normalized_query) * 1.0,
                               similarity(chunk.normalized_heading_path, :normalized_query) * 0.8,
                               similarity(chunk.normalized_content, :normalized_query) * 0.65,
                               similarity(chunk.search_text, :normalized_query) * 0.75,
                               similarity(chunk.normalized_source_uri, :normalized_query) * 0.3
                           ) * 0.3
                       ) as lexical_score
                  from knowledge_chunk chunk
                  join knowledge_index_snapshot_chunk mapping on mapping.chunk_id = chunk.id
                 where mapping.snapshot_id = :snapshot_id
                 order by lexical_score desc, chunk.chunk_index asc
                 limit :limit
                """
            ),
            {
                "snapshot_id": snapshot_id,
                "query": query,
                "normalized_query": normalized_query,
                "limit": size,
            },
        ).mappings().all()
        return {
            str(row["chunk_id"]): float(row["lexical_score"])
            for row in rows
            if row["lexical_score"] is not None and float(row["lexical_score"]) > 0
        }

    def vector_candidates(self, db: Session, snapshot_id: str, query_embedding: Sequence[float], size: int) -> dict[str, float]:
        rows = db.execute(
            text(
                """
                select chunk.id as chunk_id,
                       greatest(0.0, 1 - (chunk.embedding <=> cast(:query_embedding as vector))) as vector_score
                  from knowledge_chunk chunk
                  join knowledge_index_snapshot_chunk mapping on mapping.chunk_id = chunk.id
                 where mapping.snapshot_id = :snapshot_id
                   and chunk.embedding is not null
                 order by chunk.embedding <=> cast(:query_embedding as vector), chunk.chunk_index asc
                 limit :limit
                """
            ),
            {
                "snapshot_id": snapshot_id,
                "query_embedding": self._vector_literal(query_embedding),
                "limit": size,
            },
        ).mappings().all()
        return {
            str(row["chunk_id"]): float(row["vector_score"])
            for row in rows
            if row["vector_score"] is not None and float(row["vector_score"]) > 0
        }

    def search(self, db: Session, snapshot_id: str, query: str, retrieval_mode: str, size: int) -> dict[str, float]:
        mode = normalize_retrieval_mode(retrieval_mode)
        lexical_raw = self.lexical_candidates(db, snapshot_id, query, size if mode == "LEXICAL" else size * 2)
        lexical = self._normalize_scores(lexical_raw)
        if mode == "LEXICAL":
            return lexical

        query_embedding = embedding_client.embed_texts([query])[0]
        vector_raw = self.vector_candidates(db, snapshot_id, query_embedding, size if mode == "VECTOR" else size * 2)
        vector = self._normalize_scores(vector_raw)
        if mode == "VECTOR":
            return vector

        combined: dict[str, float] = {}
        for chunk_id in set(lexical) | set(vector):
            combined[chunk_id] = round((lexical.get(chunk_id, 0.0) * 0.45) + (vector.get(chunk_id, 0.0) * 0.55), 4)
        return dict(sorted(combined.items(), key=lambda item: item[1], reverse=True)[:size])


storage = Storage(STORAGE_SETTINGS)
embedding_client = EmbeddingClient()
retrieval_store = PostgresRetrievalStore(EMBEDDING_DIMENSIONS, EMBEDDING_MODEL)


def startup() -> None:
    internal_auth_token()
    initialize_postgres_schema()


@asynccontextmanager
async def lifespan(_: FastAPI) -> AsyncIterator[None]:
    startup()
    yield


app = FastAPI(
    title="Lynxus Knowledge Service",
    version="2.0.0",
    lifespan=lifespan,
)


@app.get("/healthz")
async def healthz() -> JSONResponse:
    def check_database() -> dict[str, object]:
        with engine.connect() as connection:
            value = connection.execute(text("select 1")).scalar_one()
        if value != 1:
            raise RuntimeError("postgres readiness probe returned unexpected result")
        return {"backend": "postgresql"}

    def check_storage() -> dict[str, object]:
        return storage.check()

    report = await build_readiness_report(
        service_name="lynxus-knowledge-service",
        instance_id=INSTANCE_ID,
        checks=[
            ReadinessCheck(name="database", probe=check_database),
            ReadinessCheck(name="storage", probe=check_storage),
        ],
    )
    report.body["embedding"] = {
        "status": "configured",
        "baseUrl": EMBEDDING_BASE_URL,
        "model": EMBEDDING_MODEL,
    }
    return JSONResponse(status_code=report.status_code, content=report.body)


@app.middleware("http")
async def inject_log_context(request: FastAPIRequest, call_next):
    if request.url.path.startswith("/internal/"):
        try:
            require_internal_bearer(request.headers.get("Authorization"))
        except HTTPException as exc:
            return JSONResponse(status_code=exc.status_code, content={"detail": exc.detail})
    traceparent = bind_request_log_context(request.headers)
    try:
        response = await call_next(request)
    finally:
        clear_log_context()
    response.headers[TRACEPARENT_HEADER] = traceparent
    return response


def get_db() -> Generator[Session, None, None]:
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


def ensure_postgres_configuration() -> None:
    try:
        url = make_url(DATABASE_URL)
    except Exception as exc:  # pragma: no cover
        raise RuntimeError(f"invalid knowledge database url: {exc}") from exc
    if url.get_backend_name() != "postgresql":
        raise RuntimeError("knowledge service requires PostgreSQL with pgvector and pg_trgm extensions")


def initialize_postgres_schema() -> None:
    ensure_postgres_configuration()
    embedding_client.validate_configuration()
    with engine.begin() as connection:
        connection.execute(text("create extension if not exists vector"))
        connection.execute(text("create extension if not exists pg_trgm"))
        Base.metadata.create_all(bind=connection)
        connection.execute(
            text("create index if not exists idx_knowledge_chunk_search_vector on knowledge_chunk using gin (search_vector)")
        )
        connection.execute(
            text("create index if not exists idx_knowledge_chunk_search_text_trgm on knowledge_chunk using gin (search_text gin_trgm_ops)")
        )
        connection.execute(
            text("create index if not exists idx_knowledge_chunk_embedding_hnsw on knowledge_chunk using hnsw (embedding vector_cosine_ops)")
        )


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().lower()


def tokenize(value: str) -> List[str]:
    normalized = normalize_text(value)
    if not normalized:
        return []
    tokens = re.findall(r"[a-z0-9_]+|[\u4e00-\u9fff]", normalized)
    if len(tokens) <= 1:
        bigrams = [normalized[index : index + 2] for index in range(max(0, len(normalized) - 1))]
        return bigrams or [normalized]
    return tokens


def estimate_tokens(value: str) -> int:
    return max(1, len(tokenize(value)))


def split_sentences(value: str) -> List[str]:
    parts = re.split(r"(?<=[。！？.!?])\s+|\n+", value)
    return [part.strip() for part in parts if part.strip()]


def detect_document_type(file_name: str, content_type: Optional[str] = None) -> str:
    suffix = Path(file_name).suffix.lower().lstrip(".")
    if suffix in SUPPORTED_FILE_TYPES:
        return suffix
    normalized_content_type = (content_type or "").lower()
    if "markdown" in normalized_content_type:
        return "md"
    if "html" in normalized_content_type:
        return "html"
    if "csv" in normalized_content_type:
        return "csv"
    if "plain" in normalized_content_type or "text/" in normalized_content_type:
        return "txt"
    if "pdf" in normalized_content_type:
        return "pdf"
    if "wordprocessingml" in normalized_content_type or "msword" in normalized_content_type:
        return "docx"
    return suffix or "txt"


def ensure_supported_document_type(file_name: str, content_type: Optional[str]) -> str:
    document_type = detect_document_type(file_name, content_type)
    if document_type not in SUPPORTED_FILE_TYPES:
        raise ValueError(f"unsupported file type: {document_type}")
    return document_type


def html_to_text_payload(raw_html: str) -> tuple[str, str]:
    html = re.sub(r"<script.*?>.*?</script>", " ", raw_html, flags=re.IGNORECASE | re.DOTALL)
    html = re.sub(r"<style.*?>.*?</style>", " ", html, flags=re.IGNORECASE | re.DOTALL)
    html = re.sub(r"(?i)<br\\s*/?>", "\n", html)
    html = re.sub(r"(?i)</p>", "\n\n", html)
    html = re.sub(r"(?i)</div>", "\n", html)
    html = re.sub(r"(?i)</li>", "\n", html)
    html = re.sub(r"(?i)</tr>", "\n", html)
    for level in range(6, 0, -1):
        pattern = rf"(?is)<h{level}[^>]*>(.*?)</h{level}>"
        html = re.sub(pattern, lambda match: f"\n\n{'#' * level} {strip_html_inline(match.group(1))}\n\n", html)
    html = re.sub(r"(?is)<title[^>]*>(.*?)</title>", lambda match: f"\n\n# {strip_html_inline(match.group(1))}\n\n", html)
    text = strip_html_preserving_breaks(html)
    title_match = re.search(r"(?is)<title[^>]*>(.*?)</title>", raw_html)
    if not title_match:
        title_match = re.search(r"(?is)<h1[^>]*>(.*?)</h1>", raw_html)
    title = strip_html_inline(title_match.group(1)) if title_match else ""
    return title, text


def strip_html_inline(value: str) -> str:
    value = re.sub(r"(?is)<[^>]+>", " ", value)
    value = value.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
    return re.sub(r"\s+", " ", value).strip()


def strip_html_preserving_breaks(value: str) -> str:
    value = re.sub(r"(?is)<[^>]+>", " ", value)
    value = value.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
    value = re.sub(r"\r\n?", "\n", value)
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r" *\n *", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def extract_pdf_text(payload: bytes) -> str:
    if PdfReader is None:
        raise ValueError("pdf parsing dependency is unavailable")
    reader = PdfReader(io.BytesIO(payload))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()


def extract_text(file_name: str, payload: bytes, content_type: Optional[str]) -> tuple[str, str]:
    document_type = ensure_supported_document_type(file_name, content_type)
    if document_type in {"txt", "md"}:
        return document_type, payload.decode("utf-8", errors="ignore")
    if document_type == "html":
        _, text = html_to_text_payload(payload.decode("utf-8", errors="ignore"))
        return document_type, text
    if document_type == "csv":
        return document_type, payload.decode("utf-8", errors="ignore")
    if document_type == "pdf":
        return document_type, extract_pdf_text(payload)
    if document_type == "docx":
        if DocxDocument is None:
            raise ValueError("docx parsing dependency is unavailable")
        document = DocxDocument(io.BytesIO(payload))
        return document_type, "\n\n".join(paragraph.text for paragraph in document.paragraphs)
    raise ValueError(f"unsupported file type: {document_type}")


def parse_document(file_record: KnowledgeFileRecord, payload: bytes) -> ParsedDocument:
    document_type = ensure_supported_document_type(file_record.file_name, file_record.content_type)
    if document_type == "html":
        raw_html = payload.decode("utf-8", errors="ignore")
        title, text = html_to_text_payload(raw_html)
        segments = parse_markdown_segments(file_record.file_name, text)
        return ParsedDocument(
            title=title or Path(file_record.file_name).stem,
            document_type=document_type,
            body_text=text,
            segments=segments,
        )
    source_uri = infer_source_uri(file_record)
    text_type, text = extract_text(file_record.file_name, payload, file_record.content_type)
    segments = parse_segments(file_record.file_name, text, source_uri)
    resolved_title = next((segment.title for segment in segments if segment.title and segment.title.strip()), Path(file_record.file_name).stem)
    return ParsedDocument(
        title=resolved_title,
        document_type=text_type,
        body_text=text,
        segments=segments,
    )


def parse_segments(file_name: str, text: str, source_uri: str) -> List[ParsedSegment]:
    document_type = detect_document_type(file_name)
    if document_type == "md":
        return parse_markdown_segments(file_name, text)
    if document_type == "csv":
        return parse_csv_segments(file_name, text)
    if document_type == "html":
        return parse_markdown_segments(file_name, text)
    return parse_plain_segments(file_name, text, source_uri)


def parse_markdown_segments(file_name: str, text: str) -> List[ParsedSegment]:
    heading_stack: List[str] = []
    segments: List[ParsedSegment] = []
    current: List[str] = []

    def flush() -> None:
        content = "\n".join(current).strip()
        if not content:
            current.clear()
            return
        title = heading_stack[-1] if heading_stack else Path(file_name).stem
        segments.extend(faq_or_plain_segments(title, " / ".join(heading_stack), content))
        current.clear()

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if is_markdown_table_separator(line):
            current.append(line)
            continue
        if not line.strip():
            current.append("")
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            flush()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            heading_stack[:] = heading_stack[: level - 1]
            heading_stack.append(title)
            continue
        current.append(line)
    flush()
    return segments or [ParsedSegment(Path(file_name).stem, "", text.strip())]


def is_markdown_table_separator(line: str) -> bool:
    candidate = line.strip().replace("|", "").replace(":", "").replace("-", "")
    return not candidate and "|" in line and "-" in line


def parse_csv_segments(file_name: str, text: str) -> List[ParsedSegment]:
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return []
    header = [cell.strip() or f"column_{index + 1}" for index, cell in enumerate(rows[0])]
    segments: List[ParsedSegment] = []
    chunk_rows: List[str] = []
    chunk_index = 0
    for row_index, row in enumerate(rows[1:], start=1):
        if not any(cell.strip() for cell in row):
            continue
        row_text = " | ".join(
            f"{header[column]}: {row[column].strip()}"
            for column in range(min(len(header), len(row)))
        )
        chunk_rows.append(row_text)
        if len(chunk_rows) >= 5:
            chunk_index += 1
            segments.append(
                ParsedSegment(
                    Path(file_name).stem,
                    f"table-row-block-{chunk_index}",
                    "\n".join([f"表头: {' | '.join(header)}", *chunk_rows]),
                )
            )
            chunk_rows = []
    if chunk_rows:
        chunk_index += 1
        segments.append(
            ParsedSegment(
                Path(file_name).stem,
                f"table-row-block-{chunk_index}",
                "\n".join([f"表头: {' | '.join(header)}", *chunk_rows]),
            )
        )
    return segments


def parse_plain_segments(file_name: str, text: str, source_uri: str) -> List[ParsedSegment]:
    paragraphs = [block.strip() for block in re.split(r"\n\s*\n", text) if block.strip()]
    if source_uri.startswith("http"):
        paragraphs = paragraphs[:200]
    segments: List[ParsedSegment] = []
    for paragraph in paragraphs:
        segments.extend(faq_or_plain_segments(Path(file_name).stem, "", paragraph))
    return segments


def faq_or_plain_segments(title: str, heading_path: str, text: str) -> List[ParsedSegment]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    segments: List[ParsedSegment] = []
    index = 0
    while index < len(lines):
        question_match = re.match(r"^(q|问)[:：]\s*(.+)$", lines[index], flags=re.IGNORECASE)
        if question_match and index + 1 < len(lines):
            answer_match = re.match(r"^(a|答)[:：]\s*(.+)$", lines[index + 1], flags=re.IGNORECASE)
            if answer_match:
                question = question_match.group(2).strip()
                answer = answer_match.group(2).strip()
                segments.append(ParsedSegment(question or title, heading_path, f"Q: {question}\nA: {answer}"))
                index += 2
                continue
        segments.extend(split_long_segment(ParsedSegment(title, heading_path, lines[index])))
        index += 1
    if segments:
        return segments
    return split_long_segment(ParsedSegment(title, heading_path, text.strip()))


def split_long_segment(segment: ParsedSegment) -> List[ParsedSegment]:
    if estimate_tokens(segment.content) <= 600:
        return [segment]
    sentences = split_sentences(segment.content)
    chunks: List[ParsedSegment] = []
    buffer: List[str] = []
    for sentence in sentences:
        candidate = "\n".join(buffer + [sentence]).strip()
        if buffer and estimate_tokens(candidate) > 400:
            chunks.append(
                ParsedSegment(
                    segment.title,
                    segment.heading_path,
                    "\n".join(buffer).strip(),
                    segment.page_number,
                )
            )
            overlap_sentences = split_sentences("\n".join(buffer))[-2:]
            buffer = overlap_sentences[:]
        buffer.append(sentence)
    if buffer:
        chunks.append(
            ParsedSegment(
                segment.title,
                segment.heading_path,
                "\n".join(buffer).strip(),
                segment.page_number,
            )
        )
    return chunks


def build_chunks(knowledge_base_id: str, document_id: str, title: str, source_uri: str, segments: List[ParsedSegment]) -> List[KnowledgeChunkRecord]:
    chunks: List[KnowledgeChunkRecord] = []
    for index, segment in enumerate(segments):
        if not segment.content.strip():
            continue
        resolved_title = segment.title or title
        normalized_title = normalize_text(resolved_title)
        normalized_heading_path = normalize_text(segment.heading_path)
        normalized_content = normalize_text(segment.content)
        normalized_source_uri = normalize_text(source_uri)
        chunks.append(
            KnowledgeChunkRecord(
                id=f"kb-chunk-{uuid.uuid4().hex[:12]}",
                knowledge_base_id=knowledge_base_id,
                document_id=document_id,
                chunk_index=index,
                title=resolved_title,
                heading_path=segment.heading_path,
                source_uri=source_uri,
                page_number=segment.page_number,
                content=segment.content.strip(),
                token_count=estimate_tokens(segment.content),
                normalized_title=normalized_title,
                normalized_heading_path=normalized_heading_path,
                normalized_content=normalized_content,
                normalized_source_uri=normalized_source_uri,
                search_text=" ".join(
                    item
                    for item in [normalized_title, normalized_heading_path, normalized_content, normalized_source_uri]
                    if item
                ),
            )
        )
    return chunks


def embedding_input_for_chunk(chunk: KnowledgeChunkRecord) -> str:
    return "\n".join(part for part in [chunk.title, chunk.heading_path, chunk.content] if part).strip()


def snippet_for_query(content: str, query_terms: List[str]) -> str:
    normalized = normalize_text(content)
    index = min((normalized.find(term) for term in query_terms if term and normalized.find(term) >= 0), default=-1)
    if index < 0:
        return content[:220]
    start = max(0, index - 60)
    end = min(len(content), start + 220)
    return content[start:end]


def normalize_retrieval_mode(value: Optional[str]) -> str:
    normalized = (value or DEFAULT_SNAPSHOT_RETRIEVAL_MODE).strip().upper()
    if normalized not in SNAPSHOT_RETRIEVAL_MODES:
        return "HYBRID"
    return normalized


def infer_source_uri(file_record: KnowledgeFileRecord) -> str:
    if file_record.source_uri:
        return file_record.source_uri
    if re.match(r"^https?://", file_record.file_name):
        return file_record.file_name
    return f"upload://{file_record.knowledge_base_id}/{file_record.file_name}"


def file_response(record: KnowledgeFileRecord) -> KnowledgeFileResponse:
    return KnowledgeFileResponse(
        id=record.id,
        knowledgeBaseId=record.knowledge_base_id,
        uploadSessionId=record.upload_session_id,
        sourceType=record.source_type,
        sourceUri=infer_source_uri(record),
        fileName=record.file_name,
        contentType=record.content_type,
        sizeBytes=record.size_bytes,
        status=record.status,
        errorMessage=record.error_message,
        createdAt=record.created_at,
        updatedAt=record.updated_at,
    )


def import_job_response(record: KnowledgeImportJobRecord) -> ImportJobResponse:
    file_record = record.knowledge_file
    return ImportJobResponse(
        id=record.id,
        knowledgeBaseId=record.knowledge_base_id,
        fileId=record.file_id,
        sourceType=file_record.source_type,
        sourceUri=infer_source_uri(file_record),
        fileName=file_record.file_name,
        status=record.status,
        stage=record.stage,
        progressPercent=record.progress_percent,
        retryCount=record.retry_count,
        retryable=record.retryable,
        failureReason=record.failure_reason,
        startedAt=record.started_at,
        createdAt=record.created_at,
        updatedAt=record.updated_at,
        completedAt=record.completed_at,
    )


def document_response(db: Session, record: KnowledgeDocumentRecord) -> DocumentResponse:
    chunk_count = int(
        db.scalar(select(func.count()).select_from(KnowledgeChunkRecord).where(KnowledgeChunkRecord.document_id == record.id)) or 0
    )
    return DocumentResponse(
        id=record.id,
        knowledgeBaseId=record.knowledge_base_id,
        fileId=record.file_id,
        title=record.title,
        sourceUri=record.source_uri,
        documentType=record.document_type,
        status=record.status,
        chunkCount=chunk_count,
        createdAt=record.created_at,
        updatedAt=record.updated_at,
    )


def chunk_count_for_document(db: Session, document_id: str) -> int:
    return int(
        db.scalar(select(func.count()).select_from(KnowledgeChunkRecord).where(KnowledgeChunkRecord.document_id == document_id)) or 0
    )


def resolve_document_for_deletion(db: Session, knowledge_base_id: str, document_id: str) -> tuple[KnowledgeDocumentRecord, KnowledgeFileRecord]:
    document = db.get(KnowledgeDocumentRecord, document_id)
    if document is None or document.knowledge_base_id != knowledge_base_id:
        raise HTTPException(status_code=404, detail="knowledge document not found")
    file_record = db.get(KnowledgeFileRecord, document.file_id)
    if file_record is None or file_record.knowledge_base_id != knowledge_base_id:
        raise HTTPException(status_code=404, detail="knowledge file not found")
    return document, file_record


def document_deletion_blockers(
    db: Session,
    knowledge_base_id: str,
    document_id: str,
) -> List[DocumentDeletionBlockerResponse]:
    snapshots = db.scalars(
        select(IndexSnapshotRecord)
        .where(IndexSnapshotRecord.knowledge_base_id == knowledge_base_id)
        .order_by(IndexSnapshotRecord.created_at.desc())
    ).all()
    if not snapshots:
        return []

    snapshot_ids = [snapshot.id for snapshot in snapshots]
    selections = db.scalars(
        select(SnapshotDocumentSelectionRecord).where(SnapshotDocumentSelectionRecord.snapshot_id.in_(snapshot_ids))
    ).all()
    selected_by_snapshot: dict[str, set[str]] = {}
    for selection in selections:
        selected_by_snapshot.setdefault(selection.snapshot_id, set()).add(selection.document_id)

    mapped_snapshot_ids = set(
        db.scalars(
            select(IndexSnapshotChunkRecord.snapshot_id).where(IndexSnapshotChunkRecord.document_id == document_id)
        ).all()
    )

    blockers: List[DocumentDeletionBlockerResponse] = []
    for snapshot in snapshots:
        selected_ids = selected_by_snapshot.get(snapshot.id)
        reason: Optional[str] = None
        if selected_ids is None:
            reason = "snapshot targets all ready documents in this knowledge base"
        elif document_id in selected_ids:
            reason = "snapshot explicitly selected this document"
        elif snapshot.id in mapped_snapshot_ids:
            reason = "snapshot already indexed chunks from this document"
        if reason is None:
            continue
        blockers.append(
            DocumentDeletionBlockerResponse(
                snapshotId=snapshot.id,
                status=snapshot.status,
                stage=snapshot.stage,
                retrievalMode=snapshot.retrieval_mode,
                reason=reason,
            )
        )
    return blockers


def build_document_deletion_preview(
    db: Session,
    knowledge_base_id: str,
    document_id: str,
) -> DocumentDeletionPreviewResponse:
    document, file_record = resolve_document_for_deletion(db, knowledge_base_id, document_id)
    blockers = document_deletion_blockers(db, knowledge_base_id, document.id)
    return DocumentDeletionPreviewResponse(
        documentId=document.id,
        knowledgeBaseId=document.knowledge_base_id,
        fileId=file_record.id,
        fileName=file_record.file_name,
        sourceUri=document.source_uri,
        title=document.title,
        chunkCount=chunk_count_for_document(db, document.id),
        canDelete=not blockers,
        blockers=blockers,
    )


def delete_document_source(
    db: Session,
    knowledge_base_id: str,
    document_id: str,
) -> DocumentDeletionResponse:
    document, file_record = resolve_document_for_deletion(db, knowledge_base_id, document_id)
    blockers = document_deletion_blockers(db, knowledge_base_id, document.id)
    if blockers:
        raise HTTPException(status_code=409, detail="knowledge document is referenced by snapshots")

    related_documents = db.scalars(
        select(KnowledgeDocumentRecord).where(KnowledgeDocumentRecord.file_id == file_record.id)
    ).all()
    related_document_ids = [item.id for item in related_documents]
    chunk_count = int(
        db.scalar(
            select(func.count())
            .select_from(KnowledgeChunkRecord)
            .where(KnowledgeChunkRecord.document_id.in_(related_document_ids or [""]))
        ) or 0
    )
    import_job_count = int(
        db.scalar(
            select(func.count())
            .select_from(KnowledgeImportJobRecord)
            .where(KnowledgeImportJobRecord.file_id == file_record.id)
        ) or 0
    )
    document_count = len(related_documents)
    upload_session_id = file_record.upload_session_id
    object_key = file_record.object_key

    if related_document_ids:
        db.query(IndexSnapshotChunkRecord).filter(
            IndexSnapshotChunkRecord.document_id.in_(related_document_ids)
        ).delete(synchronize_session=False)
        db.query(SnapshotDocumentSelectionRecord).filter(
            SnapshotDocumentSelectionRecord.document_id.in_(related_document_ids)
        ).delete(synchronize_session=False)
        db.query(KnowledgeChunkRecord).filter(
            KnowledgeChunkRecord.document_id.in_(related_document_ids)
        ).delete(synchronize_session=False)
        db.query(KnowledgeDocumentRecord).filter(
            KnowledgeDocumentRecord.id.in_(related_document_ids)
        ).delete(synchronize_session=False)

    db.query(KnowledgeImportJobRecord).filter(
        KnowledgeImportJobRecord.file_id == file_record.id
    ).delete(synchronize_session=False)
    db.query(KnowledgeFileRecord).filter(
        KnowledgeFileRecord.id == file_record.id
    ).delete(synchronize_session=False)

    remaining_files_for_session = int(
        db.scalar(
            select(func.count())
            .select_from(KnowledgeFileRecord)
            .where(KnowledgeFileRecord.upload_session_id == upload_session_id)
        ) or 0
    )
    if remaining_files_for_session == 0:
        db.query(UploadSessionRecord).filter(
            UploadSessionRecord.id == upload_session_id
        ).delete(synchronize_session=False)

    db.commit()
    deleted_storage_object = storage.delete_object(object_key)
    return DocumentDeletionResponse(
        documentId=document.id,
        knowledgeBaseId=document.knowledge_base_id,
        fileId=file_record.id,
        fileName=file_record.file_name,
        title=document.title,
        deletedChunkCount=chunk_count,
        deletedImportJobCount=import_job_count,
        deletedDocumentCount=document_count,
        deletedStorageObject=deleted_storage_object,
    )


def snapshot_response(record: IndexSnapshotRecord) -> IndexSnapshotResponse:
    return IndexSnapshotResponse(
        id=record.id,
        knowledgeBaseId=record.knowledge_base_id,
        retrievalBackend=record.retrieval_backend,
        retrievalMode=record.retrieval_mode,
        status=record.status,
        stage=record.stage,
        progressPercent=record.progress_percent,
        retryCount=record.retry_count,
        retryable=record.retryable,
        documentCount=record.document_count,
        chunkCount=record.chunk_count,
        failureReason=record.failure_reason,
        startedAt=record.started_at,
        builtAt=record.built_at,
        createdAt=record.created_at,
        updatedAt=record.updated_at,
    )

def create_import_job_attempt(db: Session, file_record: KnowledgeFileRecord, retry_count: int) -> KnowledgeImportJobRecord:
    import_job = KnowledgeImportJobRecord(
        id=f"kb-import-{uuid.uuid4().hex[:12]}",
        knowledge_base_id=file_record.knowledge_base_id,
        file_id=file_record.id,
        status="QUEUED",
        stage="QUEUED",
        progress_percent=0,
        retry_count=retry_count,
        retryable=False,
    )
    file_record.status = "UPLOADED"
    file_record.error_message = None
    file_record.updated_at = now_utc()
    db.add(import_job)
    return import_job


def sanitize_object_name(value: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9._-]+", "-", value).strip("-")
    return cleaned or "document.bin"


def infer_initial_url_file_name(url: str, title: Optional[str]) -> str:
    parsed = urlparse(url)
    leaf = Path(parsed.path).name or parsed.netloc or "page"
    leaf = leaf.split("?")[0]
    base = title.strip() if title and title.strip() else leaf
    base = sanitize_object_name(base)
    return base or "page.txt"


def infer_url_file_name(url: str, title: Optional[str], content_type: str) -> str:
    parsed = urlparse(url)
    leaf = Path(parsed.path).name or parsed.netloc or "page"
    leaf = leaf.split("?")[0]
    base = title.strip() if title and title.strip() else leaf
    base = sanitize_object_name(base)
    suffix = Path(base).suffix.lower()
    if suffix in {".md", ".txt", ".html", ".csv", ".pdf", ".docx"}:
        return base
    if "html" in content_type.lower():
        return f"{base}.html"
    if "markdown" in content_type.lower():
        return f"{base}.md"
    if "csv" in content_type.lower():
        return f"{base}.csv"
    if "pdf" in content_type.lower():
        return f"{base}.pdf"
    return f"{base}.txt"


def fetch_url_payload(url: str) -> tuple[bytes, str]:
    request = UrlRequest(url, headers={"User-Agent": URL_IMPORT_USER_AGENT})
    try:
        with urlopen(request, timeout=URL_IMPORT_TIMEOUT_SECONDS) as response:
            content_type = response.headers.get("Content-Type", "text/html")
            payload = response.read()
            return payload, content_type
    except HTTPError as exc:
        raise HTTPException(status_code=400, detail=f"url import failed with http {exc.code}: {url}") from exc
    except URLError as exc:
        raise HTTPException(status_code=400, detail=f"url import failed: {exc.reason}") from exc


def update_import_job(
    import_job: KnowledgeImportJobRecord,
    *,
    status: Optional[str] = None,
    stage: Optional[str] = None,
    progress_percent: Optional[int] = None,
    failure_reason: Optional[str] = None,
    retryable: Optional[bool] = None,
    started_at: Optional[datetime] = None,
    completed_at: Optional[datetime] = None,
) -> None:
    if status is not None:
        import_job.status = status
    if stage is not None:
        import_job.stage = stage
    if progress_percent is not None:
        import_job.progress_percent = max(0, min(100, progress_percent))
    if failure_reason is not None:
        import_job.failure_reason = failure_reason
    if retryable is not None:
        import_job.retryable = retryable
    if started_at is not None:
        import_job.started_at = started_at
    if completed_at is not None:
        import_job.completed_at = completed_at
    import_job.updated_at = now_utc()


def update_snapshot(
    snapshot: IndexSnapshotRecord,
    *,
    status: Optional[str] = None,
    stage: Optional[str] = None,
    progress_percent: Optional[int] = None,
    failure_reason: Optional[str] = None,
    retryable: Optional[bool] = None,
    started_at: Optional[datetime] = None,
    built_at: Optional[datetime] = None,
) -> None:
    if status is not None:
        snapshot.status = status
    if stage is not None:
        snapshot.stage = stage
    if progress_percent is not None:
        snapshot.progress_percent = max(0, min(100, progress_percent))
    if failure_reason is not None:
        snapshot.failure_reason = failure_reason
    if retryable is not None:
        snapshot.retryable = retryable
    if started_at is not None:
        snapshot.started_at = started_at
    if built_at is not None:
        snapshot.built_at = built_at
    snapshot.updated_at = now_utc()

@app.post("/internal/upload-sessions", response_model=UploadSessionResponse)
def create_upload_session(request: CreateUploadSessionRequest, db: Session = Depends(get_db)) -> UploadSessionResponse:
    logger.info("create upload session knowledgeBaseId=%s", request.knowledgeBaseId)
    session_record = UploadSessionRecord(
        id=f"upload-session-{uuid.uuid4().hex[:10]}",
        knowledge_base_id=request.knowledgeBaseId,
        status="OPEN",
    )
    db.add(session_record)
    db.commit()
    return UploadSessionResponse(
        id=session_record.id,
        knowledgeBaseId=session_record.knowledge_base_id,
        status=session_record.status,
        acceptedTypes=SUPPORTED_FILE_TYPES,
    )


@app.post("/internal/uploads")
def complete_upload(request: CompleteUploadRequest, db: Session = Depends(get_db)) -> dict:
    logger.info("complete upload knowledgeBaseId=%s uploadSessionId=%s fileName=%s", request.knowledgeBaseId, request.uploadSessionId, request.fileName)
    upload_session = db.get(UploadSessionRecord, request.uploadSessionId)
    if upload_session is None or upload_session.knowledge_base_id != request.knowledgeBaseId:
        raise HTTPException(status_code=404, detail="upload session not found")

    payload = base64.b64decode(request.contentBase64)
    try:
        ensure_supported_document_type(request.fileName, request.contentType)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    upload_session.status = "COMPLETED"
    upload_session.updated_at = now_utc()

    object_key = f"{request.knowledgeBaseId}/{uuid.uuid4().hex[:12]}-{sanitize_object_name(request.fileName)}"
    storage.put_bytes(object_key, payload, request.contentType)
    file_record = KnowledgeFileRecord(
        id=f"kb-file-{uuid.uuid4().hex[:12]}",
        knowledge_base_id=request.knowledgeBaseId,
        upload_session_id=request.uploadSessionId,
        source_type="FILE_UPLOAD",
        source_uri=f"upload://{request.knowledgeBaseId}/{request.fileName}",
        file_name=request.fileName,
        content_type=request.contentType or "application/octet-stream",
        object_key=object_key,
        size_bytes=len(payload),
        status="UPLOADED",
    )
    db.add(file_record)
    import_job = create_import_job_attempt(db, file_record, 0)
    db.commit()
    return {
        "file": file_response(file_record).model_dump(mode="json"),
        "importJob": import_job_response(import_job).model_dump(mode="json"),
    }


@app.post("/internal/url-imports")
def create_url_import(request: CreateUrlImportRequest, db: Session = Depends(get_db)) -> dict:
    logger.info("create url import knowledgeBaseId=%s sourceUri=%s", request.knowledgeBaseId, request.url)
    upload_session = UploadSessionRecord(
        id=f"upload-session-{uuid.uuid4().hex[:10]}",
        knowledge_base_id=request.knowledgeBaseId,
        status="COMPLETED",
    )
    file_record = KnowledgeFileRecord(
        id=f"kb-file-{uuid.uuid4().hex[:12]}",
        knowledge_base_id=request.knowledgeBaseId,
        upload_session_id=upload_session.id,
        source_type="URL",
        source_uri=request.url,
        file_name=infer_initial_url_file_name(request.url, request.title),
        content_type="application/octet-stream",
        object_key="",
        size_bytes=0,
        status="UPLOADED",
    )
    db.add(upload_session)
    db.flush()
    db.add(file_record)
    import_job = create_import_job_attempt(db, file_record, 0)
    db.commit()
    return {
        "file": file_response(file_record).model_dump(mode="json"),
        "importJob": import_job_response(import_job).model_dump(mode="json"),
    }


@app.get("/internal/knowledge-bases/{knowledge_base_id}/files", response_model=List[KnowledgeFileResponse])
def list_files(knowledge_base_id: str, db: Session = Depends(get_db)) -> List[KnowledgeFileResponse]:
    records = db.scalars(
        select(KnowledgeFileRecord)
        .where(KnowledgeFileRecord.knowledge_base_id == knowledge_base_id)
        .order_by(KnowledgeFileRecord.created_at.desc())
    ).all()
    return [file_response(record) for record in records]


@app.get("/internal/knowledge-bases/{knowledge_base_id}/import-jobs", response_model=List[ImportJobResponse])
def list_import_jobs(knowledge_base_id: str, db: Session = Depends(get_db)) -> List[ImportJobResponse]:
    records = db.scalars(
        select(KnowledgeImportJobRecord)
        .where(KnowledgeImportJobRecord.knowledge_base_id == knowledge_base_id)
        .order_by(KnowledgeImportJobRecord.created_at.desc())
    ).all()
    return [import_job_response(record) for record in records]


@app.post("/internal/import-jobs/{job_id}/retry", response_model=ImportJobResponse)
def retry_import_job(job_id: str, db: Session = Depends(get_db)) -> ImportJobResponse:
    failed_job = db.get(KnowledgeImportJobRecord, job_id)
    if failed_job is None:
        raise HTTPException(status_code=404, detail="import job not found")
    if failed_job.status != "FAILED":
        raise HTTPException(status_code=409, detail="only failed import jobs can be retried")
    file_record = db.get(KnowledgeFileRecord, failed_job.file_id)
    if file_record is None:
        raise HTTPException(status_code=404, detail="knowledge file not found")
    retry_job = create_import_job_attempt(db, file_record, failed_job.retry_count + 1)
    db.commit()
    return import_job_response(retry_job)


@app.get("/internal/knowledge-bases/{knowledge_base_id}/documents", response_model=List[DocumentResponse])
def list_documents(knowledge_base_id: str, db: Session = Depends(get_db)) -> List[DocumentResponse]:
    records = db.scalars(
        select(KnowledgeDocumentRecord)
        .where(KnowledgeDocumentRecord.knowledge_base_id == knowledge_base_id)
        .order_by(KnowledgeDocumentRecord.created_at.desc())
    ).all()
    return [document_response(db, record) for record in records]


@app.get(
    "/internal/knowledge-bases/{knowledge_base_id}/documents/{document_id}/deletion-preview",
    response_model=DocumentDeletionPreviewResponse,
)
def preview_document_deletion(knowledge_base_id: str, document_id: str, db: Session = Depends(get_db)) -> DocumentDeletionPreviewResponse:
    return build_document_deletion_preview(db, knowledge_base_id, document_id)


@app.delete(
    "/internal/knowledge-bases/{knowledge_base_id}/documents/{document_id}",
    response_model=DocumentDeletionResponse,
)
def delete_document(knowledge_base_id: str, document_id: str, db: Session = Depends(get_db)) -> DocumentDeletionResponse:
    logger.info("delete knowledge document knowledgeBaseId=%s documentId=%s", knowledge_base_id, document_id)
    return delete_document_source(db, knowledge_base_id, document_id)


@app.post("/internal/knowledge-bases/{knowledge_base_id}/index-snapshots", response_model=IndexSnapshotResponse)
def create_index_snapshot(knowledge_base_id: str, request: CreateIndexSnapshotRequest, db: Session = Depends(get_db)) -> IndexSnapshotResponse:
    snapshot = IndexSnapshotRecord(
        id=f"snapshot-{uuid.uuid4().hex[:12]}",
        knowledge_base_id=knowledge_base_id,
        retrieval_backend=DEFAULT_SNAPSHOT_RETRIEVAL_BACKEND,
        retrieval_mode=normalize_retrieval_mode(request.retrievalMode),
        status="QUEUED",
        stage="QUEUED",
        progress_percent=0,
        retry_count=0,
        retryable=False,
    )
    db.add(snapshot)
    db.flush()
    document_ids = list(dict.fromkeys(request.documentIds or []))
    if document_ids:
        db.add_all(
            [
                SnapshotDocumentSelectionRecord(
                    id=f"snapshot-document-{uuid.uuid4().hex[:12]}",
                    snapshot_id=snapshot.id,
                    document_id=document_id,
                )
                for document_id in document_ids
            ]
        )
    db.commit()
    return snapshot_response(snapshot)


@app.get("/internal/knowledge-bases/{knowledge_base_id}/index-snapshots", response_model=List[IndexSnapshotResponse])
def list_index_snapshots(knowledge_base_id: str, db: Session = Depends(get_db)) -> List[IndexSnapshotResponse]:
    records = db.scalars(
        select(IndexSnapshotRecord)
        .where(IndexSnapshotRecord.knowledge_base_id == knowledge_base_id)
        .order_by(IndexSnapshotRecord.created_at.desc())
    ).all()
    return [snapshot_response(record) for record in records]


@app.post("/internal/index-snapshots/{snapshot_id}/retry", response_model=IndexSnapshotResponse)
def retry_index_snapshot(snapshot_id: str, db: Session = Depends(get_db)) -> IndexSnapshotResponse:
    failed_snapshot = db.get(IndexSnapshotRecord, snapshot_id)
    if failed_snapshot is None:
        raise HTTPException(status_code=404, detail="index snapshot not found")
    if failed_snapshot.status != "FAILED":
        raise HTTPException(status_code=409, detail="only failed index snapshots can be retried")
    retry_snapshot = IndexSnapshotRecord(
        id=f"snapshot-{uuid.uuid4().hex[:12]}",
        knowledge_base_id=failed_snapshot.knowledge_base_id,
        retrieval_backend=failed_snapshot.retrieval_backend,
        retrieval_mode=failed_snapshot.retrieval_mode,
        status="QUEUED",
        stage="QUEUED",
        progress_percent=0,
        retry_count=failed_snapshot.retry_count + 1,
        retryable=False,
    )
    db.add(retry_snapshot)
    selected_document_ids = [
        selection.document_id
        for selection in db.scalars(
            select(SnapshotDocumentSelectionRecord).where(SnapshotDocumentSelectionRecord.snapshot_id == failed_snapshot.id)
        ).all()
    ]
    if selected_document_ids:
        db.add_all(
            [
                SnapshotDocumentSelectionRecord(
                    id=f"snapshot-document-{uuid.uuid4().hex[:12]}",
                    snapshot_id=retry_snapshot.id,
                    document_id=document_id,
                )
                for document_id in selected_document_ids
            ]
        )
    db.commit()
    return snapshot_response(retry_snapshot)


@app.post("/internal/import-jobs/{job_id}/run", response_model=ImportJobResponse)
def run_import_job(job_id: str, db: Session = Depends(get_db)) -> ImportJobResponse:
    logger.info("run import job importJobId=%s", job_id)
    import_job = db.get(KnowledgeImportJobRecord, job_id)
    if import_job is None:
        raise HTTPException(status_code=404, detail="import job not found")
    if import_job.status != "QUEUED":
        raise HTTPException(status_code=409, detail="only queued import jobs can be run")
    file_record = db.get(KnowledgeFileRecord, import_job.file_id)
    if file_record is None:
        raise HTTPException(status_code=404, detail="knowledge file not found")

    import_job_id = import_job.id
    file_id = file_record.id
    update_import_job(import_job, status="RUNNING", stage="FETCHING_SOURCE", progress_percent=10, retryable=False, started_at=import_job.started_at or now_utc())
    file_record.status = "IMPORTING"
    file_record.updated_at = now_utc()
    db.commit()

    try:
        if file_record.source_type == "URL":
            try:
                payload, content_type = fetch_url_payload(file_record.source_uri)
            except HTTPException as exc:
                raise ValueError(exc.detail) from exc
            resolved_file_name = infer_url_file_name(file_record.source_uri, file_record.file_name, content_type)
            ensure_supported_document_type(resolved_file_name, content_type)
            object_key = f"{file_record.knowledge_base_id}/{uuid.uuid4().hex[:12]}-{sanitize_object_name(resolved_file_name)}"
            storage.put_bytes(object_key, payload, content_type)
            file_record.file_name = resolved_file_name
            file_record.content_type = content_type or "application/octet-stream"
            file_record.object_key = object_key
            file_record.size_bytes = len(payload)
        else:
            try:
                payload = storage.get_bytes(file_record.object_key)
            except Exception as exc:
                raise ValueError(f"source fetch failed: {exc}") from exc

        update_import_job(import_job, stage="PARSING", progress_percent=35)
        db.commit()
        parsed = parse_document(file_record, payload)
        if not parsed.body_text.strip():
            raise ValueError("document parsing failed: no extractable text found")

        existing_documents = db.scalars(
            select(KnowledgeDocumentRecord).where(KnowledgeDocumentRecord.file_id == file_record.id)
        ).all()
        for existing_document in existing_documents:
            db.query(KnowledgeChunkRecord).filter(KnowledgeChunkRecord.document_id == existing_document.id).delete()
            db.delete(existing_document)
        db.flush()

        update_import_job(import_job, stage="CHUNKING", progress_percent=65)
        db.commit()
        document = KnowledgeDocumentRecord(
            id=f"kb-document-{uuid.uuid4().hex[:12]}",
            knowledge_base_id=file_record.knowledge_base_id,
            file_id=file_record.id,
            title=parsed.title,
            source_uri=infer_source_uri(file_record),
            document_type=parsed.document_type,
            status="READY",
            body_text=parsed.body_text,
        )
        chunks = build_chunks(file_record.knowledge_base_id, document.id, document.title, document.source_uri, parsed.segments)
        if not chunks:
            raise ValueError("chunk generation failed: no chunks generated from document")

        update_import_job(import_job, stage="PERSISTING", progress_percent=90)
        db.add(document)
        db.add_all(chunks)
        db.flush()
        retrieval_store.refresh_chunk_search_vectors(db, [chunk.id for chunk in chunks])
        embeddings = embedding_client.embed_texts([embedding_input_for_chunk(chunk) for chunk in chunks])
        if len(embeddings) != len(chunks):
            raise ValueError("embedding provider returned mismatched chunk embeddings")
        db.execute(
            text(
                """
                update knowledge_chunk
                   set embedding = cast(:embedding as vector),
                       embedding_model = :embedding_model,
                       embedding_dimensions = :embedding_dimensions
                 where id = :chunk_id
                """
            ),
            [
                {
                    "chunk_id": chunk.id,
                    "embedding": retrieval_store._vector_literal(embedding),
                    "embedding_model": EMBEDDING_MODEL,
                    "embedding_dimensions": EMBEDDING_DIMENSIONS,
                }
                for chunk, embedding in zip(chunks, embeddings)
            ],
        )
        import_job.failure_reason = None
        update_import_job(
            import_job,
            status="SUCCEEDED",
            stage="SUCCEEDED",
            progress_percent=100,
            retryable=False,
            completed_at=now_utc(),
        )
        file_record.status = "IMPORTED"
        file_record.error_message = None
        file_record.updated_at = now_utc()
        db.commit()
    except Exception as exc:
        db.rollback()
        import_job = db.get(KnowledgeImportJobRecord, import_job_id)
        file_record = db.get(KnowledgeFileRecord, file_id)
        if import_job is None or file_record is None:
            raise
        update_import_job(
            import_job,
            status="FAILED",
            stage="FAILED",
            failure_reason=str(exc),
            retryable=True,
            completed_at=now_utc(),
        )
        file_record.status = "FAILED"
        file_record.error_message = str(exc)
        file_record.updated_at = now_utc()
        db.commit()
    logger.info("import job completed importJobId=%s status=%s stage=%s", job_id, import_job.status, import_job.stage)
    return import_job_response(import_job)


@app.post("/internal/index-snapshots/{snapshot_id}/build", response_model=IndexSnapshotResponse)
def build_index_snapshot(snapshot_id: str, db: Session = Depends(get_db)) -> IndexSnapshotResponse:
    logger.info("build index snapshot snapshotId=%s", snapshot_id)
    snapshot = db.get(IndexSnapshotRecord, snapshot_id)
    if snapshot is None:
        raise HTTPException(status_code=404, detail="index snapshot not found")
    if snapshot.status != "QUEUED":
        raise HTTPException(status_code=409, detail="only queued index snapshots can be built")

    update_snapshot(snapshot, status="RUNNING", stage="COLLECTING_DOCUMENTS", progress_percent=10, retryable=False, started_at=snapshot.started_at or now_utc())
    db.query(IndexSnapshotChunkRecord).filter(IndexSnapshotChunkRecord.snapshot_id == snapshot_id).delete()
    db.commit()

    try:
        selected_document_ids = [
            selection.document_id
            for selection in db.scalars(
                select(SnapshotDocumentSelectionRecord).where(SnapshotDocumentSelectionRecord.snapshot_id == snapshot.id)
            ).all()
        ]
        query = select(KnowledgeDocumentRecord).where(
            KnowledgeDocumentRecord.knowledge_base_id == snapshot.knowledge_base_id,
            KnowledgeDocumentRecord.status == "READY",
        )
        documents = db.scalars(query).all()
        if selected_document_ids:
            documents = [document for document in documents if document.id in set(selected_document_ids)]
        if not documents:
            raise ValueError("no ready documents available for snapshot")

        update_snapshot(snapshot, stage="INDEXING", progress_percent=60)
        db.commit()
        chunk_records: List[IndexSnapshotChunkRecord] = []
        chunk_count = 0
        for document in documents:
            chunks = db.scalars(
                select(KnowledgeChunkRecord)
                .where(KnowledgeChunkRecord.document_id == document.id)
                .order_by(KnowledgeChunkRecord.chunk_index.asc())
            ).all()
            for chunk in chunks:
                chunk_records.append(
                    IndexSnapshotChunkRecord(
                        id=f"snapshot-chunk-{uuid.uuid4().hex[:12]}",
                        snapshot_id=snapshot.id,
                        chunk_id=chunk.id,
                        document_id=document.id,
                    )
                )
            chunk_count += len(chunks)

        if chunk_count == 0:
            raise ValueError("no chunks available for snapshot")

        db.add_all(chunk_records)
        db.flush()
        retrieval_store.ensure_snapshot_chunks_compatible(db, snapshot.id)
        snapshot.failure_reason = None
        update_snapshot(
            snapshot,
            status="READY",
            stage="READY",
            progress_percent=100,
            retryable=False,
            built_at=now_utc(),
        )
        snapshot.document_count = len(documents)
        snapshot.chunk_count = chunk_count
        snapshot.retrieval_backend = DEFAULT_SNAPSHOT_RETRIEVAL_BACKEND
        db.commit()
    except Exception as exc:
        db.rollback()
        snapshot = db.get(IndexSnapshotRecord, snapshot_id)
        if snapshot is None:
            raise
        update_snapshot(
            snapshot,
            status="FAILED",
            stage="FAILED",
            failure_reason=str(exc),
            retryable=True,
        )
        snapshot.document_count = 0
        snapshot.chunk_count = 0
        db.commit()
    logger.info("index snapshot completed snapshotId=%s status=%s stage=%s", snapshot_id, snapshot.status, snapshot.stage)
    return snapshot_response(snapshot)


@app.get("/internal/index-snapshots/{snapshot_id}", response_model=IndexSnapshotResponse)
def get_index_snapshot(snapshot_id: str, db: Session = Depends(get_db)) -> IndexSnapshotResponse:
    snapshot = db.get(IndexSnapshotRecord, snapshot_id)
    if snapshot is None:
        raise HTTPException(status_code=404, detail="index snapshot not found")
    return snapshot_response(snapshot)


@app.post("/internal/retrieve", response_model=RetrieveResponse)
def retrieve(request: RetrieveRequest, db: Session = Depends(get_db)) -> RetrieveResponse:
    logger.info("retrieve knowledge snapshotId=%s topK=%s", request.indexSnapshotId, request.topK)
    if not request.query or not request.query.strip():
        raise HTTPException(status_code=400, detail="query is required")
    snapshot = db.get(IndexSnapshotRecord, request.indexSnapshotId)
    if snapshot is None or snapshot.status != "READY":
        raise HTTPException(status_code=404, detail="index snapshot is not ready")
    if snapshot.retrieval_backend != DEFAULT_SNAPSHOT_RETRIEVAL_BACKEND:
        raise HTTPException(status_code=500, detail=f"unsupported retrieval backend: {snapshot.retrieval_backend}")
    retrieval_mode = normalize_retrieval_mode(request.retrievalMode or snapshot.retrieval_mode)
    top_k = max(1, min(request.topK, MAX_RERANK_CANDIDATES))
    min_score = max(0.0, request.minScore)
    candidate_scores = retrieval_store.search(
        db,
        snapshot.id,
        request.query.strip(),
        retrieval_mode,
        max(MAX_RERANK_CANDIDATES, top_k * LEXICAL_CANDIDATE_MULTIPLIER),
    )
    if not candidate_scores:
        return RetrieveResponse(hits=[], lowConfidence=True)

    mappings = db.scalars(select(IndexSnapshotChunkRecord).where(IndexSnapshotChunkRecord.snapshot_id == snapshot.id)).all()
    mapping_index = {mapping.chunk_id: mapping for mapping in mappings}
    candidates: List[RetrieveHit] = []
    for chunk_id, score in candidate_scores.items():
        mapping = mapping_index.get(chunk_id)
        if mapping is None:
            continue
        chunk = db.get(KnowledgeChunkRecord, mapping.chunk_id)
        document = db.get(KnowledgeDocumentRecord, mapping.document_id)
        if chunk is None or document is None:
            continue
        if score < min_score:
            continue
        candidates.append(
            RetrieveHit(
                chunkId=chunk.id,
                documentId=document.id,
                documentTitle=document.title,
                sourceUri=chunk.source_uri,
                snippet=snippet_for_query(chunk.content, tokenize(request.query.strip())),
                score=round(score, 4),
                pageNumber=chunk.page_number,
                headingPath=chunk.heading_path,
            )
        )
    candidates.sort(key=lambda item: item.score, reverse=True)
    hits = candidates[:top_k]
    logger.info("retrieve knowledge completed snapshotId=%s hitCount=%s", request.indexSnapshotId, len(hits))
    return RetrieveResponse(hits=hits, lowConfidence=not hits)


@app.post("/internal/read-chunks", response_model=ReadChunksResponse)
def read_chunks(request: ReadChunksRequest, db: Session = Depends(get_db)) -> ReadChunksResponse:
    logger.info("read knowledge chunks snapshotId=%s chunkCount=%s", request.indexSnapshotId, len(request.chunkIds or []))
    snapshot = db.get(IndexSnapshotRecord, request.indexSnapshotId)
    if snapshot is None or snapshot.status != "READY":
        raise HTTPException(status_code=404, detail="index snapshot is not ready")
    requested_chunk_ids = []
    for item in request.chunkIds or []:
        chunk_id = str(item).strip()
        if chunk_id and chunk_id not in requested_chunk_ids:
            requested_chunk_ids.append(chunk_id)
    if not requested_chunk_ids:
        return ReadChunksResponse(chunks=[])

    mappings = db.scalars(select(IndexSnapshotChunkRecord).where(IndexSnapshotChunkRecord.snapshot_id == snapshot.id)).all()
    mapping_index = {mapping.chunk_id: mapping for mapping in mappings}
    chunks: List[ReadChunk] = []
    for chunk_id in requested_chunk_ids:
        mapping = mapping_index.get(chunk_id)
        if mapping is None:
            continue
        chunk = db.get(KnowledgeChunkRecord, mapping.chunk_id)
        document = db.get(KnowledgeDocumentRecord, mapping.document_id)
        if chunk is None or document is None:
            continue
        chunks.append(
            ReadChunk(
                chunkId=chunk.id,
                documentId=document.id,
                documentTitle=document.title,
                sourceUri=chunk.source_uri,
                headingPath=chunk.heading_path,
                pageNumber=chunk.page_number,
                content=chunk.content,
            )
        )
    logger.info("read knowledge chunks completed snapshotId=%s returnedCount=%s", request.indexSnapshotId, len(chunks))
    return ReadChunksResponse(chunks=chunks)
